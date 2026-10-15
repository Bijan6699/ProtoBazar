# -*- coding: utf-8 -*-
"""Pricing Streamlit page; calculation and workbook-preservation helpers retained."""
from __future__ import annotations
import io
import re
import zipfile
try:
    from defusedxml import ElementTree as ET
except ImportError:  # install defusedxml in production; fallback preserves local compatibility
    import xml.etree.ElementTree as ET
from datetime import date, datetime
from pathlib import Path
from typing import Any
import numpy as np
import pandas as pd
import openpyxl
import streamlit as st
"""ماژول قیمت‌گذاری پروتو بازار: تحلیل، پیشنهاد قیمت و خروجی اکسل با حفظ ساختار فایل اصلی.
اجرا: streamlit run pricing_app.py
"""
LOGO_B64 = "UklGRmJMAABXRUJQVlA4IFZMAABQLAGdASqAAoACPjEYh0OiIYmkrBABgllbuF2sRswb9JQr8F+E/2H+F9E/j3q48Ufev0J/fP/j/k/uj/rdNvWvlK+ZfrH+h/uv+Y/4/96////7++P+w/3/sn/pP+M/7f+T+Ab+Pfzf/H/3f/Df8z+1//////XH6xv3N9Sf9P/wn/i/xX7//LR/u/+3/sPe3/bv8n/2f9d/e/kS/oP+G/7v5//G57N/oW/1X/Z//b2dv+x+43w+/tj/8f9X7mv+A/5v5//+3w/+kHjP+Ofuf+P/uv7R/171d/HPsf79/fP2X/vPQ2iQfIPt/+S/w/7s+23f38yv9z1Bfxf+Uf4D8qP8V+43LUgD/Pf7T/pv8V/kv/f/jPU8+jPYP7T/9j3Bf1w/3nloeLXQF/p/+C/83+l92/+1/93+l/2frd/PP8t/4v858C/8y/sf+2/wn+Q/+f+p///LdiK7h8tKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0OOOyyh7nkYYE4FPcdZ8sdkzGz9B/Ykx+eRhd+yOOVZEE+aAaowwJwKe47Sh7nkWtP0L9Vrd4IDVd8HteSqU8YBL7GC3ium5ZGEzTfwXIWdr+xSECnuO0oe55GGBOBTWvujuXVIBiTe1tJpMK3fw4DJJEV8q5IhM+SSdzBSNRhJiKW7KlWux+1odlj8G/3XKAQKe47Sh7nkYYCLZvOcoC8LMiUfxarvLWWnAr6/aJgXLaLm6PFM6ML9kPSd7G9a7L3qGLSaHAp7jtKHueRhgI2LsAy1UqJUg4kzaY1J2X53TwUM7UlumCaMQcKSF93h2lgNdwwjwTDYgzjcdMxb+N9yvxBhI2pNj9bwpvk00KUVXd1laDq6/Nxj0+4fLSh7nkWyqsm7Y//2cP//4AS3lOxT//82I+ioOuQ8tlkNj6Fq0a6USiumuLt3OYpCs4TsLYmXNLya0DuJGWPlpQ9zyMMAtUQbAJ4lIjG+tMQHU9srb1b5hHJjBRwSBQ3YvWvMllJvxGurgSrPZQYHlPDjDCHZN8RHY+kAwpr+PSVVyw2/6Epi9F1ZCMYRhzbbFpkxVB8tKHueRhA+/icSBDpJ4i0FEsLQqk0g63m+e8lzK99i8JVn8xUYnEd9rbqGc1/7Gn2MBGhzue0iWxaTWaXeZ0y7C73X/iWWp5DdFJFNqujl9ujsLpQcs934P2ACtqVoHTnMw5gyuHplam87F6SiLA9zyMMCVqmVV9NfC2tzvt/oTuF7z7vzP0ho6M1atLMAWM/MCMHbAUWv3w7YlmYgi6rwe2apH7O9lSlMryiegXLc9VAxT/BPHP+w9s+vbdCBH7L66VHjwafRRxuutNo5wyh+v59/m1KN/3lN7LVpq+1IIkFoPZvuHy0F6L4QcZKZFelPeYujtJRbbPzG88aCY0Pr0wCbaIPrR5OiMJFih08s1l4vzyZN6hgC8FsjiWjFvFPEPQO+Ou7p+ncyFm5lGPzEfUCJ309+jKK+tmw7ivta6//BQgTkeF4ZN3FuRFKmaOnMv4aXppD/w7uHn10CbRIkmV3eOEV3OCTJEJ1OBWbcr5l5ltSs77h8tBUuYOcfu31k4kbkyhRO9mHukMyVcnK8vwWNIv4v0pnN+bb1GFDGgOnN/BkophEgW4x1XlXyX+RSZ58EBrhcJWBQ7ZhcenCW9sOZCJ4GqGmmhBu6e2RMzZsy/J0HRDiDa+mVqll+ewdX9izugGAVnCLQS5zUEiEBPZzyMHLd/vebzLqXuM34obyk0qFlulUTNwg/hCqx+a8i2q2UmBO3JD1eOLh85+Okmqnh82An2OtmAnWQlm52aspOqZkXf5ZeEVH7Uq3RfikAPf7yrICRAJOg6/Pwtunv9BwW2UFu++FP5aTqiHOCg2RJoJeGAQKe403V3I53I9+/sLBhx/UyOA5bg9zfHEe/YDBzUcZq8fSpfOW8FTwW+RlioUWd8v42QIVkI99QnV57oARRDJS4C+MYHtFKeuDgOPUx1lP/sWcyWvX9L4D0/BoxVp3rMRbp9SHCIau1EyK/X4j9Nhdn02guv+tBHvHaSIUFaPo0DTh8tKFRTTIq/pUHXtzFWzRc4ceLbHRRv8C15M1bXzYaJlRVSv4f9/m3WxJ007n//oWb2vaJvmhIMXFSMW87Kjyie1izO8kFeuzwUH+8/967ni3Ca8fhv8PTN4JtCjLVX4E4GTRJeYbRMoYyIXM1cH4Omt8THEqZMPU+ljhc+WlD4PwLTd+u32yOAMbrOlTjS++bJiKeiqj8yhRyHl5vPX14iMudNazG7s0X8fixuOKwmhgeIgq5R56BTUzk9lX+dGJvMYRo1pEBHrrt4e6COT9pgv/fTUNglYCYx/Hrf29lJek2M5/PvI+WQFU5S5WvjBQVICnuOutm5LYl3UCJ4eaLYreSuE4bKYKts88M5JLlveiGP7t/YVxVdPvdXSFursgrbvv/VEjtkizNhVJM2hsB0Nr/SbihGCu5bv+51Xhu05Wy3co4tLvgIhs57K0iqvMf5yXDx+QbPcfAp7jtKUfT8+kgRqoYHnolEEmq8iDJDweRldIs9KenQBZY7jj8Ez06fOIRsacVGoQdG5bIhWupVYGvU748j3zL8TR/RpL/oZwkYT7okyT3SjJ5N/mcCnuO0oe55H74ujXoV5Xh3VrKf5LwQpH2LBj/aAvPSlADbXVKR1XlkIGZLFkxFtIP4YE4FPcdpQ9zyMMCcCnuueYLyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPcdpQ9zyMMCcCnuO0oe55GGBOBT3HaUPc8jDAnAp7jtKHueRhgTgU9x2lD3PIwwJwKe47Sh7nkYYE4FPbsAAD+/+8CAAAAAAAAAAAAAAAC0xO0skEotg4h9Clx7SLKzwp61591kGinF/N7heNqEj4dtDBL+zb7EVoqlLBgFUYmXVMUCRsGckshkOzJPyy3XE6PqX1K8f+mYgKoghLThY/TGiCUSSZRMNDslTYBJN23G0TIZZktq7qGd4N1Nvfsm2frqHDZyOnCHjRobzhHTZ82u7puhFA2NlY5hvxIjxhcmqyU1lPG1xWwFeecM3QxdsAj0OVD7jIqe673yap9+5nVpDhBwp5bOlOv/UTHjfSsUtoP5pvteKR8s1b9haAAE30UfsH7owLKK6T9YjQJNqYO3Es72U3+/TjahQzcDw9yoErsms0ozCKN2YEHA31rayXNKGJwPdDxuRl6rPRHTgR34kbVQCskGQ8IdTAxqqeJ5SHj5ZjjgPFQCG14Zy7t5rJPpwP2hzuxiHzOMWilMJ6EpR7zrHmSBbhUgTxxa82y3xkMmsmSxy1vV3bEcpvBGWT1PtPtT0PKsKRvLGJKso/QmeakL2AOEZr/ttqNS6ZzcLspst/xztvDVY+2PlYVOYg26PH7uFGlo8f9JkgMnfO17bbs7KKglVbaGJBcyCJ5VDcq7ILq3FUm2hw/+XYTrwWpAA9JG4vfD5r6kTStniSr0hl1Y7HhbGRk7vdVX5VPYT5vNG7XLXF9JK+ZonIU8n1EMVnutPjwB/jF9DzVVUx6/VRy1qUAq67++i9P776KKJlNuWzWFMDqQtqbMN9LfkNou5kQh8rvdiTAojQ/CJOsNTqG7NFgI02EoyZjMHzcCYzxs00nPVQpO2zYGdlMsgAAYlR+uRSl3v4gTd4hwwY3+nKe+f+PAS5sxcGJ/LaQR9hJAg994PCrNM+XggTA6oFPr780Ig1bR7gJ4/G8uOam1q/vBruCUiWnCOtZDLVIEssk40w5do9eZjBSvsUk0jAftmrV8d9+PN27HelPzqKtKXl5/5M68CMk/Ib0LVY+PXM2soo89M1uT9aI4GkeOr0tqqJnFl9YEW1Ot2oAZOUmP/e4eti1X9nSlR4LmBZs51Ih6+eilAK20Cmb3qbblpfFVQqq1oXg11oFSGsgxs4Nr/Y8txGiESLMRN3NGL3f6z+51RafEtDXPJbNpj9fHjzPc0UFaEzvh+LyOUjZtui3RxYpAkSkh9auuRx6RpZCBUVeRCNuPf6+P6ZANZ08HQrYbp0wlvGp8Li0AAFWDc2ZVQlR2phu99dMotGmod66zd3ZmwSy/fI+sPIsS7zUCE1mEhr6ExeSrcsckE7QMtHE5jOH+ag5P31WYEAkdX+c7hD1Hi5jqzR50NVlv8jO5stzr7k9PPXaR89A11HrLENVQ+u1tfvCa0GYLJEsjb1S3fD3ORTErPS8kbPIvKShioUMjZawbiqLeuNMGpuJTi/UQcaoIZ3xVRMeMFfA+u5CpMVKXXivhVBmTUzGbo/WHiURN2xhx5HwVnGpWW5jeGoK0DKwd5g+StXP9UYHPzl7hWB8It64/bHGAAAAKKmMG6qCzn1Pm/yRsW/9X4bPNB2ScffW6U8+ZRTDCxbeua6L4K1JM/lip0kxexnOMdrH/9g7mDnysg0tenqDfr/W1tIuocoiyZ37fxnTz+Bnj5KwBalEhCSVxOZ6806I6f+FJ3zu+By+JjncMHrD2HiAv5ue3ql5Y/Aoi9ERpl/IYvD/LXXLy3c5S0fkeo1tgVNg9AhFgZ6x8oYCL4vx0g38Y9OJCvEb6M7LUwIscKUIGak+NyU/AH7B4k60Ppzk2KWm6Go2gEnbTzyZtvUP1DQ8mIXm8jyZwJwynTRYE8LNSH7klQP03WI+ox2fAnKUtCExiivzgKi7CG2aB3aO0NY81zTjNYe+M/39Z/69j132mr5BoAvuvynzakLuI76jr3Tr+CQfOi2ZdX2hjqIIDLcV09kmy0aeP+8TMepwh/uz33PTLFS9Hl9AOf6wWO3fS32Dm+ssnlr8JhvqPn/IQSAyvw7wrNsmd8v1/f/luzt3HYhd1mfk5p6zT2GX9XP2eSZsGp/9PWS2YBFPD3JhCU5HdZefbiTUQc/HBiAuh2FXwzg79kIhw4KtO33pSm7Nx4g+/MfCyV4p6uTnk9ItKjInZiKrtKdyviHWEVdyFQygloVG2+TKz5SU6/H6+XimdOcNMaEW2K8jTPlxVz+6rCzUbXan6cI5aDT7X6HWofda9s+Jy0KD8PtyY6xApCz3GZv0pwdqRbA5+l/nOzHiTvd77wpgekcUjFdwcdIf7m/wnm5CWdGwk1PTTOrUAgwVU1oF9UeJ8MgblH8biSXdPa5mVorFCiTNBUS3Bbuu0u8q1suDJJ21nbJkOLx3Ww8/mWmQpfcyDsGu+eYVi0RO54wsDntppj/zQo1asJ5IzUaJPTVty93fVOy8UxnCzFwNpFzI22y+7SU/1ZkGWHq8Q9F6oKI4JU1AykECYVWigm8srhbbkuzI6pUkPdM8WGNG4gPiKSfUmOvfECUas0sTRK0ext94sFUnJS069eqbRgfRnZcPStRe8Ra14O3GXzbl3Ml0KErWcUv+rgPOan9bx/8igZaE19qEXCOKMnq5BaCnU7//dM2RJy3AFJyaxD6NJ3Ryi4qJ/S+0PoNiRu45OzIFoufiSjmWAT3CGjFfo3FHfhugRMEiwKhVmLqBiql+ZzJ9Tb6udR2Qx73VpfXWro5YnYGx+RoR1U0Ca9jF3WspBdLyclwma5NyTEcyyAoi5GO25Xn3jreE5QUG0O2tD5ahZJ28XDp9vHJ6nlDv2LfGn5BzQZ4WuRnQawtZvDkkRtJvQDoe3txMigjxp6pYbiRF2eK96GG0TqY3meI3QKFsiSuQDuo6eHl7TeuAKW48rYLePmIkTMiKFKcEfs5JVOs+1afOLlWELg0GCZMumMw47jS9HRXr/BLPPwzw14NFLxSEO++8twtnUO4n7UceTCxYwfUx2pnv7bViIzHIQE5p/xRzwtOCh7nNw52NSDQbC6uucCoMjrPEJbw7Hy7hc8mChZh4OsOWii5V6yo9t/plhcxvbumHBp4Y6GFxfWOhFycNfUD8UgsWNi70aunRNR73dHIb7czQk9bYESHawnOiuoNZrfQV/5A3o2XoOtORzoUezpx1wZysWNG8TuyQCchP7jMyk5HaKyPvL6DArPTdNWETU3y9PA8VSf9PzyK0Jl9JDO+uzGetzzZHGECDyPusxcG1tlcwseIKT7f/1bHu9GgdJJPK/YKPNYZkm3slze9ivdwBllF8Fudp0wA7gMzr68x2K1uBNkgNdJJNN0GjcQPVKoG4qX4I2snU6et/X8ELl06d9rA6yAbGMQRPGYg/tb0n/A/6d/q61xaW4JDKnoAgcqvtU+lv3Abqca8ipItSC2RFSHt30vuLKE66jKj8FAXT9TL39wJaV6K01lSoyowXQN0boZ6dIk0bv5+Avy8bkYUKgogSzLYYAFM8qMTp/5cpp1YRj05Ui7ICRxB+/X0uol/xUa+8mWaAOfgn7XhbmcFJsc/zzKd2UjugR+um6nvxTrE/RyRpWD14kAyKCN1z8o0eUwaGITAAjCWM/xCtqMpLI4rgqABER1NAbf30K+j8CcmKYKY7dzd0nz8ii9uMSJ8Dbc5l06gVgZpNR2+PcVhnFf9mPEp7JV6HcGTvyRtdpxhwkWpJ1c60jYvMWIp5g5uBh0sUQ4Wffyb6B0lZqm94kQNVPV8nTJ49ktw4UQxBzlGHcX9G0h0X7qxjGMHn/TxwB9STuoFDUILTmu2wvh2IE48Ljzltv3oha7myWHdwjug6nzEruI1lC2i0j9aiXGGMOj8rOW9E46AVibElqGR29WkUohjLey3yqdalgZnBjVFyf4jURfZP5nuOTb1Vfr4tvCHAEK2bmBS7vprG+hZF/xagP7mIb/AZgqXzRTjhA5veuHA+UdjeAmjePjOhGMrt3IbqRYX/jnUrxZBZDTyoA8txT0OpnwHRqh7qXPngLCatNwgcqg+RqxTr+6+jeO4Rrm4lbQBlmgd9ylmSuZy7i8L+scXOdnj7saHQh8mpUrmcT7/0fgvb050ptuSRiWeuseZ3GmnVe1fmcwc5bPKJBXxtL+T9zGIGjM5RJD1fWdXdCxPAipEYoPQ0jcyY8nGEQZm6XGVj24iCbDuRXtWHunx+zutVmPAHBtia5VKbIsCOzeAnYFIVE4p59rwL3I89olJWQYkHYsVBHuLeIKPhR70/9vU0KU5VYTxKLga34AxoTUaXIAABYmtd5V+cH1RVMRm3Ku0h/5RPFVyIcnBG1L9Ug1CyWs54Jp+TSr97dCpn0/veHvyQaQkzxdl55n/M7NAWGc9HSkXU14LXP3h6vlqwSKcEyl4EQfNSukak0KkUnbDumJVcOYTNbwYgL4XPoW+CAHkLZob14rQ0lNZgtL7xrhs8MLo3Sqk4wqQ/AIRHGf13wjNW8iWYiObK5nTHSz61hJL8gc1p9gpTgipqyNJswDljTsviDyqdE9s9waaZtqwW8UCb9PPYZndYLH5gg9v6HYjnxdh9e+LwKq7Pw7ilXzCY23q7Q9K1OD/Tw5gyzE8LT7KgQquDiZVeqvRP3+p9BnxgKWOf+V6B4L2p4RaoEUGIpSe3GGu6sXUkzbXqzbGepp3Uvz019IIn4RKqz4qm4uAOr3QWsF38E9CZTPqo+guWZVu4rweBmcJKAdY/seDyP5sE9Jx4DJgfGrtuSjPq5UbURYYm+5uu5tTWK8PkwbWcV8EwmCJb7vt2rrumU0qT1bTn3pYZE3U1JLh0s+F8L+4FWNBm/WPAz013A8dApva71ljzJ73o9SqdPY9cMtcexxOzheYCfZqj/LV9DMZt/MSkF3WRwChZRc1PHiSqbhwUx4j72mWpSQiq1jVvbG024HwgxCF3DhiPQtY/CMJW99JPIYRQs923b4RpgnqrlyxLmCDGdjFfbhelatJqvSgydbezi04Ea7v9mYyflzfoe4ixePwMCmGkTVRm/2M6J8B/y5KhyR+huC4Vzhzrll1x4UeaLP8bX7Jj0cne6h0cT/zRzTlNImWXUUgQnFGSfXIt3mrOAndSysfMT+g6/WTTDD4Tnbk1P37NMDSdL0dRuA8+YM/AFLp5z59epd5Thls2Mis4jD3cgBf6NMXumlucABWVmGarXqibI01TlEiG6wZ5qrpu74ceo/yhu0daK9yGkYbU0oyVh10NVxJJ2t9VdeheOYIu3/AzudLr3MElYq9vqrf0pCgfrZWo2ZfBeoOCJS6gApjIkuKBHyLM1b9Gw3EDpENmb9iSjTo4ZOjYBrudtDgSIGrwbQFYBQjGBj8bD0BX7tREEQRFgO1VLTYda3ghyBcSInTI89UP3E0LVXuQ9wkXnh5vOEmGMry/l3z3ll1OWqzRfyjhvttrPEvsLgUYa9qOyEJms2f9ut0ybS6YTLDipr6WqY1dPicV8Dl7sxDCKsncrqHdwkJu/H8tZvot6Lfmd8HmuuX2kC+QCoVO6XQeJGv0l8n4hPCdfK9gP7tAnJwk+XfmWvRrh1V1fRJhSXdRNHUiDP90PVFjeUbH7Q53dU2NBbLT60Wf0sGdeW7SeI7SGQnHQi+QfN0PSf9eLN9r54J4jMZowdA9+OpvJA8Hpo7k/a2Wq/9R0H1LOzJJ5macB2nsLaFSNq6lrXt3ETytI/VLFY70b/BSdD8jbi/DIIdS05K6HmVVdyEFWdc7Gj6UJ6AY5LyLzrXfzKY9BmaSm3ndFdiPO54yhGs4L8vnmlXc0jxdD5rtVnayfWu+AqfzgV0IiMk9sLpxwSVRCmUYycPfVFjxInUPFL6oNkXYDXBUF/Nx90E1hofGOJl6a4C6J72dNwAgl8xrMuoleag9cp1Z5cjMf/juMAtGd1XA24yoDmEOkQkc0LqylCmHiIOpkZy/Zdgyy5pIp5isngnFk/RM1T9nS2kwWae2kaB10DVCdlEEFQB8KNc8rNTpyYsRKMRXTMBbgNkPZ53X2nl9Py4diydSpRMYzZiPyJdL6B1d61fatoWGuW/V0PBwLU2O2Gz3zeytNn9BHyrrUXfsSZUHYxbDbnnmuy8rNgfJb6pbl8Rjs1fTTqZsSo1U+Vdu3POv0Vm6bWjvMEz+4kZZEZ9McQYqeb5+xoML95yIYJKDy2hcdQjqe8fcJUXgKzhdPmV22x4G/zd3cmq5lY4BOWNjkOhyjG2A9+Pmpa+T2RoFHthg0z5uRWIlY/85MO93EpqylnTsm0wT/DYFiYmMlOErnnc3zM8y0SvZZeZR8P7dGe/nuSdvIsYNsY1ycgly1Ht5rRCozzP3G+DNawFPL5aLE6oRA38od+o8S0KlDrtuHo+O1VSh2vIQXKL+T5+sW6EsBB/yVmHeIoAzDjXdIVD1oMtanx3NlkpzmhytN1uJ287n6iMYBIWdUsqbNKeAMtruXy7F9p08W8G3dUYckbnraNrLCNNrhlMuH4+ryEMBs9PW9nkdpou+DpXjcb61LrOTU6krCjX8H8bqWK4gTqC7rc/VmYOSHbtqU91ER6hJLswd3WfJFWY8zyQLKyD0USGuWEUZypE7Hv1/4o8T1W1S4TqLux0dlee+w8eXkVLsJrTxSMvCyk+beaxeJUprog8kjTY9VCmWcEBZNXMeU48VTe4bY0WRLSlXO+St2g6jB5KUynUs/Yo3BAeILa5gpw14YnE/auUWnQF/aUgfn7J3SOPKKvQl4m+GqlNAtaWyGlq29kooa49tJZBbJph+pvvsvQMIuKPydtvBWNuQmyoS+DIfIcImxC3tBV0CEfnGZv0SKCpbAG/XnHCLciJqWLUXZEG59b81qyi+WkpHBRtjDveKQd3te9C1QsnrAoaBvpc+atwSKjjvZs2nxJUSFDC4Kb95mZEhrGKQMgJfNk3jt1euV/TqIW9rScB/42oOx8HuK5NE4UBrejBT4L9JJSCSdO4whVHXRA+oVwtpIxtv0fi6/YqUXNGTMcNKaqqulKjwaK5xudJjrVjTmpwJdgW4N4VGTh5VkPBZzn2USNp4ZrF80t9pnJa+shl9K6CAznG4TvcXxUUiErZkR/edrLQXd+fmz6NjwX9jk4JMXbks6pJYcSP6vpmRjj0ViiRxhnhfvHN40USrPhDyvawFIE1tV7syrm4rJZTIaH7h97HAQFjSbAGbyad8c+3gaNDzeocZQIocbmXY2m1QLTkiPgTH8IS6iCKGL/qLHZ7AGOkY8AUmQtc5Ewy8MAtN7UDXjVGD4GJfg51bDagp+hrQvJCkiFkY+aOmk54ra0Dk430hM2wwtGyucMD08AtzUFe2WsD1mUZo07OhXFs7lw00NGhrzzEo+qGE64UR+D9hW3iB0MuDZJUOt4naHZR8Az+wNsNaCkp3NArRcSvexlVug2Dj7svwe6ZSJMTVgDlfK0FQVlDkHgtvSfWgoLhqSPvEIWAlalAdGGfHAYpIu4VjddptefY5iz4JwhSd97fpVKe9ebP6B6+m5CBVv5PKi+uDpVbGPS7uHfxGML4alUXvD2nhsAL61mK+57ZiKzoQV21WKtlMVnCKUy0mLlNX1pzb2SMq1Qk0OnMT92hdHKZX0+WePRx9wBvluWz0vo5ICtn74YlW3y+70yTL9r1BuBn18+F7uQgy0kJMArIJ1G6nPq+1pmLrjdLmp2D/bVBlaX3bb/QLdcDhiUcV2Y7KEhwk3lBuPeXW8MzSRCXXvH9re2s0jge0HWXgmCmKdKJ3WelYx6SYVT2TgwmlzSuHo/lzLB+sO6Yq01xJpu7P5wVJutpDFTrPkDAAf3zOPr8FO4cq8DNzbNncWeIklK7cuCek2zm7incvL2Ye//jzQkvx/KpySyo62+mLrQvpLwY9qUhmEJatgNZgkRDXJcKr/3/3saNUtKHZmTGFr8dYp4HbLIhVxhvrntbHZjKTV7SQv1U9LJB1h5VMjZwHUpnP0jVBLRAwM6tE0trLeJts+DuiswrVPebKvkwbZZbCx3TbfZD/tEMk2YBr0usjUDBbwKLgwy0Hvdjab7ILjbSTDlzjfDa2EFU4DIwLFHF3FL3UJ0A/egXNLI1V95Jx3tE4yQ3vFtp3sB9W0cd82RN/kgszCCWkYQ2hE4Zd6gXJpwEut1C+3S/CjNH1FUt9aWq1L8iZb1+R8w3WQGdImFQRWiiC/v2Quu72K24ZqW/W3OjvEb2PbL0u2JNNRhet3SZ1ZIpMTJQCr1QTZ5TOVDFuN6ASxim6BJ2vNY/7951aJAmx07jr2hD3cyGveI/qF0+XzOhJvU6ZgUOUqs/n+m14TYRWsBdgMYTWM6QfVzqSPt7sRZBZAy4+rpLOzlZPvKIn3meAAEKuWWLSYHxNu+Xd7dVcZWYtcZ8t4GLhsHcdx300rrscte33koGUYyEa3WZV98AbeHDe+cuNu4pjvMxpqbNA3Nlteg2gJW2dS4fzC8zVtK/9o8IhOxjm/KixccrpQ6K6JnugAiwPf9tC97NyFw8PKz3qTe1issVbRpZ/CWv1nNuqm0QnyjEdZ0cfQkHJv3PP+LwDwP9tVAI5mctUQKecrKiF1yiL6IGfy7uyyzjjaa0cbtkX5Ud3TlVn8RTh1m53MWIAulS9GYoOtnonidubZ95idn3Fv8QC+IcF+tI/IRm6y8Vs6OYz7p9zjDScaksWjbjKYGIA7i2ZTQqr8MLLH7VgQF8w3Z94Dk7tgdLhr7omxwVy7neyN78odBZhJ1zioQzFjlAlyMqfB86CVJO406qL56qL2WTpP8p+FJhc9Rkele2BXxWqCIZxu4bPFGQVF6xJaqz3ejHOA3LQPVYZ9quZ5JG7aM9+ve5zxyzol3vSsHgnPpStPAY8Y6sVAuKHlojSGs+98oYTYz8JxbwFhbwUTe57Pcavdtrt6WyL6RXpzOsTdYgoGkyG/mtacoYu5TnaKPErJAvqajRfc4WJeVKjKinwHtek3igZRvq/dnshxEuGWv/r3d8Hku0BvQo7+6wTdHgBehCic8zI2hx9kJpkYr2NuofP0ekWIKJPNEsaunHWnua+o73PSl45Sni7ctFRj+rmkbrw5gMRvAsbLj9yqOlMO/zvAEEq3UPF/EemCoW2qwCgctBeYIEFc05nojftM3oMegn5a5zKw4CifqKRDAn+o9vi2WBLVNR1SJjVFZXmUthdL+VKp8acN8guXIYkkJ3XtCMsvaDMd9dQ1WC0R062cthK/OmJmN5vgZ5vobBPTfFDVSY61hXhPIesG/Khbap8F/pdITQsgUIL3Zv/2zoMYx0XVnSoqN7i29C3hSJaMhjNyTYsn1ETPQcn6+FKlEbZEo6+rfFFhk4RHusSK3MlYbBX5+hqNKKQ+gtaquqnyyp+bkRnt2G0XDBawGpcvatrNJ5yIepPOOrGYOcnc2pbBeaoKF8Kr+1sUqMSr0DEBpQP+46IPPX+Xi1uMgfzgOAiK0JtunV72g4R1B30bk4CtBTyF0JPRRGjfGpJAj8463tFloOVgtq1y99ZrmZRHgO8OPXNN8MuoZcs3HbONzA1OOBoctrARBVPPoQGDYZ+3wC7uHa3cv2GrL3YRZhWe2ShLG1aym8foKKM+EiEsErD1k1YgBInBqpzckEpH0eE8eibjdd0zh1dxmEI1qrS42TYlKkxhj0Cvp9VlkzfUcOvgzV2MXQMiHGbnRbt5fwpD/hkqnIOdelhLXSEUjefcc1Y4Y/4mFDl0UeFVe2FuHWVt9f+93KbEABWepYWa6ECeyJyindhJR2TSWvNyE2oPpRHoM+a0lSDsNRgWf15AsSvJ9RYa2/r4hTBaKChDW6I70vt2eVbrVVy4AOuOlbwMMrEH0wTOAP1R/4ReotvbSHTyfniJH5qlRdVB8y5U5OkMuQc7fbgba7qJncNmB3BajKxHvsEsgjYW3f4M8cZ0ofqAFZolJoBUaE/Gmz1ZZlx+5EGuiTVUcPFTjLLSuyr/KaQeHPfugJkNy2ussUQ+n67p8cAXI9U7F+HONJ372xSWn+AbYUHISCLhHo/9ayTMJYM+Vj7jRZWGGUxnklfNsg7ISRzB8skBCWTsJCYZO67p9d4FE/c02X67W+8c6aVG9JlutFyDZfLlSolJDaTaiDrgY1ib4cQQW/84GSUiktzrGtW9vQimLayBDDG+nTRyfuCMkZOBCUcMu8+FgPTbW8yvQQp8VMtx5QXoDrx0sbtbInIyoz5wYCIXLxg/u65WXw8ODAccbfuuLXRDZTLnpeI3IyTTbi+5wA2pMsXHLfz8h3ee33nx5KmvfC8Cf9kcqBgrbvVboHiRa8f2L+m2fjZkZPoyxyawYA988Xm1mXkp96k0G8jlf7ieuI2Uxnb+SPPSbhk5Y7XPCdr6BoJaD3+CcEF6qEP425d6hXd6Pko29QkJNXKl6xJHWSKdlJnMNLrMcfIClk9olezFJdfdS0PaG83scSGsagG3n7T0D+EFTHK81hWCoNkSCs5gOxXaVh7eAbXrPRQszR5HnjUc10825FhbcyBYKUnIjRa9OIw/T/MEoXtw50DifKBj2ipMORrHZDQGz2yf8GvckyhIWRsFFL5JkV5YA0B6PCBr1S+/LvMqFaNNCGJhVAaARnTGfMa8qQdt9dcm5hTLHKfIRalNwNOnuLDbE+qUZWRB9jt6dGX1ABIG6ilpOqMO975jBwM3cDRbr+SzDoFk7FPyHyKqgQ0LWHEZD8UUMee6p0QkNvej/aflIW8fKxXUarRjHuv7d4PJ3KX6C8UoOSvF5mlYpotWUXXvhdPTVydNnCuqjiTUW07La8PToHxOUceJ35gOgetJ6K3fCmIFWccEkUT/f4wEllMDD/COGFVwkgQ18uKRaQq3KtI8wm//Bv4cc53PrbEgPqCuxEe/5qoQtbPzV2DGqYzuk5lBCgBcr0ZIPpUyg/bVBPWq5J6H14LSUXATtcTKq+xLCO0XR/inR6LHZs8s4l4VS/8oD5Sw/X6+RdaFW3LZDh/kEEk/Ag06HE7gr94XWrjxi2k+FDlnLanVPzEvJu92NtIGNPKwCa4LKRFomMZPVdBQxMtOCbNv0JdWxu2dSz9ssdUe9a/wEH6BJrZhHQiqtHU2J1jHsGAosRiAPzI0hwVO16qlD60KnYba2OjxXCjgAG1S8HAm2LIQXJvksg54bv8N3C6ki0PT4kM/ohyT8WBo9vSm+YVlwC/D9Ub/KM8MrviBGhW/oBWQDb+QVZeekXWV7etBzN2hkl3aWGLlWqjVn+LENshadguGQOV312qxK1ERvEKDv1ZkO9tvPcZWsC7yLdWkVIkeDsnxv7SY0j69WV35mUsklir5Ah2HP4jLMyI2dla8G5dEj8Pka0bYi9CcE4I/hPbGqk/Dgz/IArk20yiTef+957QRBlNt6Nb87QJmlXPltKVwWmBBKZAF02bzG3+m3JVQgui4KVaS6MBbXou6KMhGbH+CkiYBqwZ5PAEVX1zdceJ6y/8ETNdNP5Z41ahOvsTFBpbrJBj/NUqUCBpOxCcjwTg7zg+KNjJR8DvoMKj3CmdlQeNR9iuVQWrOtPOk+X8reLrZzlGNTa0UxG29m2sQTjL5YGTxIb5rCMKMEJqhLjNl9/Joewkuq1xYUj5nohPpguYf7TT+Xv92B69VfMgfVTBr3yc3D12q61HNnLoMFWw8cvfdeBXhTi3VMVePlCvb5XMhC2nb0hwYcaSoSvGlMJfm6bvaukPFGrX7fvH9zcxHKgq9RbFmeQg3BTyTaR5DYPEb6pB1UZVQ3Z+e+4H8tmmyijqGqQVChlHWit6j4M5GZsImQSS0Mw3sEbzwByis96+nXr/0RKh41nGxbxl0tB2HHSH3SwtXWfG4qGXNu1+JRD7VVPt60tFa9cUJsssCt4FL/0hmi1UhtNTJ+fx6GxMIzeUH5nLrFekXuZg56ib3st/azfmOgCD7B50f550gA2VQrwp/cTTm40x+bgQjxyy38pHk0nUHhPW5+f93rDFfCaJoY4N9rSz7Rof5hFFSm2QbuG6ebQET18YkIIOUqJFfIxJFaTfXr9Qx++UllKW+ajveeUf+1YoJt16weMmIy4KrxteaL4pYjXqgDHTZEIT8sYeCwDCinopYLNNvVqxTQzdWGzyvgf7OB6hUWpfFqfMFczjVjF8sb3pRy1DlhMEXHLFWY/kWZLbeYm2QAqiwsej9qHBPOPPv/HaQ/3/iSMOObp2yx8FKZ5QPmUjxW4UT/I4rloY7vlpbcVN5ll3TbqcYc3P4daatbmV3+cswy+FPHio5eYHjC9UhVh7r9Do8GIV0qQjd7QbU5CLbRhmiGq4dJEfWIpBU987LXzJ+8H+o/ssOPr6kokZ1Kit8bypCC4/17Vof9fE2Z7JclTUQy9w3VOKuU6y33zpYt64IrDFEtjTxmNHEAQD+zgcw33Gm4FLpqC2hTIbkMt2ee2pUvQ9cMlGcNX+zKXIODl3+reI2yM5v7V4aOcpBAgGvX/rXguO5fX0WwDUy5WtltqYwxSVNeFrN5SpBIhb3Ouw4KbmBDoSjzDhfYA0kQQ+Xxir+pTo39Z1YN4nixmWC8NvbHnT7vWtFPToVB+iwOgvOryRIYqFHKxBC4aAtIzUHxYK/PbV8MJWvXPWGyuykQOy+4F+rtJ4NBGlq6KtNaqdxDC3+hv4VsFY87IcTeFrqku3rL1f6nsrLKRY08sNjaOuEVcGkjX+7+VvmvqH96++2zchdC67RSUxkNHRVR40V7zJ0WOA0Jlv5YA57vMks4JrcnGkKWTbp1r/suzOZPote1lCNZ+Bexvyv7h/B7Uo48gCj0nmWThJPl6qbh0uf2WZdbgb9j9AckzcQN52VACqc4TAdyX1r28HNemZ9HKRLKDnlL+m16Am8x08L/ZZt2iQjI5MHqgt7rytQcjWFE+h0Nv+naZBx1xA1aw6Hq2qeP+zZVe5ZPuCsqzsaCY91+v8NvLFHtuh3b/wGeeQJlHNuO4ON4IexQYrGjN3Xe7Jb58ngt9kGbdGIugsyKgX3pKJEDvJmWML4X8v4NOV0rQctu4PKxYYKHDd1e1mRTYpJivgIOLU2g5DpEBRmQRHcZ5xnGMG4BouF5xv9VjON0Xu/5HcAx8YctaseWH4emvY8pChGIrPPMSB+ngcxwtlAFNjAmcKbgfn0Ygl9kGspgzUGcvTeKpVN1T2PU8BLuH+VJQnHqWNZpf5xUhSO0Q0QiBo1pbRnEh2l9qcajg3Wq7IsYBw/WFZWbfwxN7H7Z/D/QBZ/DUzE6w1CUYe663z8GyfyWzj314Pa64lS5+we/ukh2zjPBJ+HVhQkNoPRVPaCixXNsonCg9qeRfeAajNApYQK3yDyuu79tUTVvx5dOUs8tyHU9pAvuwXaUet1jj2o1tn4Wsp6bayK448i8TBBaEsKi8dS9gbStdy0Whvw2mZzMfdVy6ZRGF2fYT+6HiW0ZfT1uqvwYnJ+pJtNpc5Fp6N27sbNxCS7lhVg1hvNfBQxvzA6golsN3r0XhqN+OfbodfCdQ8TXTW/c+MgxJQP9jS9faf5VfUOYB1zN/vaXxrmRJBqSw7k1mlalEhbbnBzZLVemlNK5q/8L55VguQLhFS5UAfao+mzudz9ikJwmrOmhOkEKUa3FvLllUubLqJH3ZWXS5KMVBOIyB8PxZUxWuJfybC9sohhIC1rBv7+uM918VnyUEL3HlyjqKYgZ/3MKnhtA4IgAgqYnAvZz59PNsfo4wyFRI8OGa/LK0ZysZqj4Rhyl9fRpukGkVFYzdTOFTlTH/RlPXRAKJDIxuZgUn1ClfysvJt4ie9rdD8sJVTPdU4VLIR9NuPN3asyHeb2Kw5fPqxcstugRMwMqJ3ABL+TjQrAngXm9UxQpWr++8piuHGNsG+yHg6/fgmNdcCiH1v6SAPBS60gkRtV3U1GiBNE/KPHDvXQnhPDLlF/nQj8fjIosYZKwNMT7RoASpWuQT/xCQqNc6wXyTPa5gh/1rviPZshR6v6ufLh4buGx0xsfwD4GKmPKpFRxs4ln5rMWKh+MeaMiaafpMn4Hz6GPE0TRSkN8I2g7rYaUbUT48hzj9S00d/M/Bnz6f0ggIb40Zdpmj8pWNmZ+Ya3CypOukFd2oiv0NtzHACOq1hiV7KQ/zRQ1V9gVPyOeKBh5+Uf8jZdEbYjFuGa7UqMhYBybpO+CAWH3DD/cvMHqjJbISHqXQD/AqVvxunNcJlxJhWpzbOQQjF7ttbw6+HTiUOBB+0Q75K8fYTCL3AlDmPZCiZPHmwbtVRTFz14bZA3FnyPX+66T6SJOWqCTmsBg/q7pTfYyGcHMl+iFBaZ0FEHuMHIOIC7+a45Bzda17giibaTjeROAy+jdeO6FYtEjvgoaXPKPq1lCZWyQtpdrlUm/w8L5o8ysJVQzWY87pkg0sISfb1vlj+o+bzdJC7J8i4Fa+bAuyMSh2TwXr+yHCT47+3agdJHA/NbhRBlKZPEJvj1zvcd5c/4eWrGIqamsVbRQCbva68GZUzIARMbDA0JvzKBXLYuFqRGlelAYwTa/uQ1nXADKYBI8kC5GVwhnYeoA7laHkXljCqxpbyBjfVUqGEm7hEJc9oL8h5z6Yp0rX3w9IWmGPJIA6GmJStOC6MgJSnrZEGyuejokqCpuC70TjgST00BOsSYZ6p+7y1zUhGE58hdn4nMaVY0rVTmTPyP12nUKjyfod9oUkt6juvI/MKojzsFAXUwFUp+qyEkYLV7ZEn96PXfNams+9Ga0XXR10MjbB4YD5ywXxMrneUnbcjWWwa45tbF2bwdjSMd84KrLbJus4M0CIpjr1beJpF4raQbLa2kj+vl+RyYoa+mqMRfMZG/0mGttTNwYF2i9OPUb/0afDj+eCxHvBZz5VvinXxwhJckQG/WcPzMzSaV5msk2F5yFK9GSFXYAgNOi9wuup9wSM38vHLwL9iRA/WRV+T7ypnaoS9aTm+zV15mNfD0OylchXnTK2UaT+w5P2x1vVCd/aF3P42/MDZAZ6PqEJzQr802J1rqLZuO9jdfz3u0YiYh8pL+2N5Vl/fibi2FZKb/TFTnCFAoVGI58qQOBUODph2f19E9M3s2d2EhhwhBaF69V6zRFkahzfCc676Ce4PcJRVTdF1HoW8euHANRU1zn8hwg+5Z8mbULVKLWBo5WaZs2qAn4ZiOplyeI8k72GdxTzFOG8djJwupp6WUqJO6X2ed+Z14RH8Sp+Ap//Z4i9EFKwHLMgTi/+Rg9VgyD+W2GS3Y+oJxmegi027Vv2DA/UMERSS0TdTFikFcf36gvDqFQy3WxE/frDiL0OU9NleZc8TPEDxirW/rNEKZzJtt2M24S+fcXDovHh1CSXpJCvt+XTTfAErQUcXooJd1BnMgcSvrgAamsr9OFazz4QuIaH9El1IJygPTXn5CCP2xIk9DS5cPzeuN4bJbJ/GPbBCY8WXLMqfH8oiyK5Vd5rE0Vb46/XLvV0fgw82TrpIKAQqlblyXwDqRMqE7bN0j9vqHMsBxAIWb9iRFZpH+CZz6rL1kR4o95jF5LE/3aZDVW8wWEg69oNKQd5j2bBaogTJxke2vRV9OZxEfyeeCYvDZgGYox0Lr9pdvstrZjTsouOLNaITx5ZmYO66JSnRECQBt352PbYPiOtyFrAgtrbW30Yi6BPU+Ub7mxJfxNKS/ZcMWlXYFuKvUtTHncg6PCj9A9wC1nnv3D0CMtd4Xkk+2fmj4b1CFbUmgapjrndFvnM+OFrQRYLfGpj8uz2AiIu0HT3pBtizr646KaHwZt6tnczGcGKT+09IrV6EDRKaUDY1x4MLKk1nir2AsrrmoAG/qEdQzASoIhKeA9utK/yHoTy58iX+tQmaXXMCn+giRyQk+ZphcLTLcXLY+Dvzq7AN9tXHHr+sOpsMzTMghlWyIYXn6BoEAXq06ldNYtlw+AWR1MyOzmUmSdz98/oG55bIT0fa+B6Wkr85Y/33A2wqjonj0cQtH0fuLLCqlz9aHziMkVxZSf3Dmc4ZQ3RuH/GA9VOoC/0bW8Bd7LqCN7pZv88rrRytiWlNKGKV/4GUyGjEbdPQcm37Q+sCGlgOCE8yRx94dBInH7ZiV/T33riOfKKTGeaILCMqdhY7eUKfUm+pg/8un5lqqSvvjJ4A2v1uVSD6xr9+UfGUnQCZq7TddN65SXG/3owADx1Wz5opzbQPEpQIvUnV5PWlNOvPb1XpEL0YypP6+w6J5YR+pDjjQjnZ5ev5NusTUTY/yfNaxtY/ZkBIMdVdZsT/DLr0YoAWLhuRaYH6XkbAZddgZ2c9OLUwcvwUFpJaqr5Q8ltqN8SCvRTRaiKmyDOpgk2xdTIiSpcp5rPTFVNOvlgaJcauLAILvt3+OQ6vw4f3AeIBOuR9Ire72ylE+SUkWt4zJXbyzlg4xAnnmV59EH3bna9m5h9Aigs3mcn0ZcbiYap6IxrhGJNebKy7nwVqwOGnBPNFDWNbwisIUZY4PEQYJYjB61/Y2k9qiWamyU/wg4NBv/lFachW+Fs+vjGDMDKWibgrx/ahXNZ4ry6jaN1qK0zEuPNHd3p0z2IuFr4Wezl/OkrmPnb/UhlFBg3+J9HptIidwQznc8MJ1Fpxj231qxNiNYa4iWmAnrH5NGPb44zGu4jE41Xw2AJm3Pyb/JyALTinPqKzqTvThUFwELBukE8IEwzHHmzPwwyaaOGd/vYCAqqGiw8EOkov+8/S5dBg3123f/7+8M6yT16+l8Z3haK90xuOFz+ir/Fvyv4rX1FkqnSnGMMZjnVPHuFlVn2fwd2ye5moeRU1qd5p1M3YezGeP0q6kWsnyjtPVlASYdhjqJgfCtJU5SB7HnR/tlPj9XymGpNpP5MSKXHWIFXnb9RR2nD/zuaqsBX5oWfrFk5BySMMt/ziPzRHtCn2iifcqbwbe5p6fYleAhmYamCZxWOcsHnpEOBpD/+rhFRHo2HNHXJdjEdJiOjnMtz+wi6UXFk+TgoE1ilDqQHZWlnYiWkmhsu02TA8eXR50REQIZsh7CmgoUEhaP/oMzxZAmDzgHTHlLAhOG/wQhfs0ViwBJro+dwiLiJhQURlV3iGzr33vemAfw9SGlV4pfuPlwXrWSncQ2eaFcOp2cYDbY3q0/S1HUHbaQc3dLm50+mtnMBhtsx16zsJPMiRR2FVeCTYdMWV6r0HFU7YTLI4rBWdwbNT2sFZ5vNsc1eMihqkN2ZQm0uaui8TQTKxRS6t+b9ZwO9wLx71jpJqQpKiYO5FDO/3xIXBeNvwCaASKDXZKWV636WE/vTggzoSttj/buw7gyiD6rlhD2RTAFozlnJXWLGH3qZ4f6dp2PH6OmMAU5ShCLDBrqkRNJVa+QXlAJkcYwVlciScQkJxWi34FmOa9N66VlD7gnwhAz2DfOm/LIAy1A//hPHRRqGKCfgijlg0VZlw/UWMxAw4IPUL3FRuGXXebzfWIG0xiiPTuZtlO908ISdJwjx6PcJx/n49hkvNorG6ygJCdxB7lNYIuwBI0y9zc2FMq/2fTrF1EIqfpLt5HAKm4VV8O8rhyCorxWGDZCZ/A7IK1/IlfUjyqNp/MVp+F/fSqpOBwYYm3QgMgGma6q+56a3goz+Kts3cu4iFWgqPVCNn9+jT83KMneU9JXn/J7mCNaIw3pfLaUZvJCN7gOTFpaSbZGV63zNlGJq2yWeFQbVaCNlFvVIPnGgs4x0ZhRSMYMh8Bw3Kxnga9qZdF/G0mIw/r51sr65lvmdS4yjmtEgp9B9aR7EesggIBnfe2qVQVOxU+YNi4bafrTbrsD+ZnVZSamWc5pq2k/76h0QExJ0QXT4abLohuyI6KagFwFm7KBcQC4Jsg9qXRnLJcjBCm6QG1Xpo/JdF90N/dsSKdqkv5nxIr/DrQb8HfULeL5+fIMNQyiXelq+MAeSBm5ygc9gv2EX1hiJHyQM0LnD9vMNyelBYl7JjlYdYcvSmGzfLxj1sWlgC6rDBkmQjqLw+WRZw8/ppNDvgQftupaWN/8QWtZNR8hj/MLJfTcxFJavKo/ovTdTYGSo1ZNgcrS5MDRoiwQA0ABCg9UbfskgDq5nUcYwvMBWg2PGXFvWO4ITdKDx87BufErbLnDHdI48jwgsqQ3PagsFpSxBSXzkkUZVO/R1q4nkffIthOUFM1stIIYRRp+dbOPdSldo1gCFJ/tQGDMDe43Fo7uaI6gxY7X9VsXwxdzJP+leJ9irsdaW79z0xnFnByiorJ3nUTd1MaJfrErMMMtctl8DH7PNo0Szu6UPeJqrk2vrLXnCBxFgPeHicArQGctr7DnJY/+o8ZzhtUd79p+Af3SayLe+0rVvWKtv3652RmbRjz84Zg42rKegZHV0kpcvnaq629/2N71f0izjvmUThzObfzfSCmjSybW4OnqvhVzJBATu4TXOz5VZ0KcUHVcf84gcOsuoQWo5kbDImqjEzBDe4feVa/oqrrL69zAWt2VDLuZHp33BqVHOHW/CaYN0JssVTfRXqhiG7v5FvTla0HHZnp267mdY9y5fE/UfVHHfw98fROOQfZh2bFQxAEdjAivUyZLKD9UWc1Yu4LE8mju1iVrgQD2Y8hdw4ko0zW/sgAK9SHOdQmusR0qnwXKLTVdcsHDcCFY4lwQhTacySxA09UW9sjZfo62jEAWCRJiT1rj4wAZIqEMigE8EKtuHch+s2PVOTt7P517mJd3SxJa3SGKx1MgTw0fTB0SEaFN96laT1kPlwGO2p5WOMuYBHmjpxBf7Z1Eu4NKznqdvLLPLRU2495vDjcTwdTAxZei6uq44EjZ51Jk/eqZgA2gl4SdVKo+1LhfD6FedjeR9W9hguESOWFmgGPY9WhehI+ustZUQ1BU8kJgQcyFFFZtX/luZBlz2g1BMyn5EDq+o6/SCDPVDYfr1SCC6IJkkC3Kl/D/q2J4/aSQiWzbySK54njxgDR6FDSxsts36ioa4r87JyqtMENW1jB/6Zgp0gcGWQoj1LeeqZHrltmG9colAAvXfpRHAIbF8Y4Kzz2UqVI0EE0SiNDkNKUWov3kchPOpInuyWePwaWhTJobfTZ9Z0vi6x6ZsiqAnllWffZJFcYqAmQ3rE/1+1Mq+vgYC9CGs5jEve9nDwVxIZ6u27YiHLpuvUDDJP0GXydsyPOicBhVRPFNRI+NOL/2YAxGvCRbx7NKGrGMSbD/I9jeSDzyWdlMgtBWF+HJ4og5PEzMsTGWlNxU0aky3SyOpFLH1WACeAD9NyAGXS5noLqsuzt0fj5niYeELPG+TUnHCw6E/2mGpdtG8dDDF/NLY9SaTjYWvyj+/9kSqw0HG6ifKOPfk7rzmjCoSRzA10HPJiekqlcAVxR2G/ebbSdks3KRhXUB3AYkliZZW6tgvuwkCTW3HNQm3G2c/Qmw3ZTre5U8tHT+UZIRgUJyDv7KsxzieeT85gnbRdpYs613B4HGWfzgHfNCOXnLo95gsjU3G1wMn6LpyW5tLsSqMkk2ws3GaOhklICPDhXUEA4F5LAXH7xp9u65GGfkS4PB93a7yOb97ufSkHShkjKAfshVjPHwIVcvuD4m7WnbDN0jq2Itispvw2wvOH7SMLT4JeUrm/r8/l0V+7ZYIrTLoG4+Ul86pCg1hGBd3a6zPjDzJOqRooJ13wihfVxBBIuuDWHq5s0jW6hIQ7cu/9qKJKGIfQGtxNDXcDDNdp/T1JwDCNecTmL7d9U6es/VQiw5+P0y1bBkBeFiDrEPCHQO28ZhvBb6niQkabsBmfR/6OiS2bdpQ8amPTUVyCdu6Oo956Oat/fnN4TuvFy5RiYKQhGVtxPVTCKS9OLCK+ScXuqSMrEYmUuHZnTyjdqtOlPiJueUQV8o8A8H6mrgJj4dOGKfwP0ZNjkMQ2s8H/h6HqIhqG0PnADNRX6YKAC64PdDcIyUL8x1olFdWZxUyAXct2e6mVtky+iTNdRDNYFglxKU3ELZWbFJCI020trQ2ccitcnH9rnUCq9wNrpPCxevHUVnopF0mkVyqMcq6e/SgpXQRCnex1JBlo5qbGtvlbLxx5q+eXI8sODTu9sOzbBIkKzPJ7BwCgTLXvWAf7GzN9kGYu6XlRfG1SlehUIfcs9llWRxqC4Ovdo76q3YC7Gv3jYvcWDHhWeX1NXa/EevwcqxtyHZBXq+Nf3QNVjrpR+RXbjAxg1mgECb8Gv6ZAidpXzi6nH7aMA1VAxoDG7ZEhoW2UihPSIoYD1wyR+ddPYSmHOTic/N+f6eWD8MsLBI9GPjSP94G6jLqjyAp0JbLFF81NxIMi6vL01FglcNx+KXjjS4QPYVDdPjse4WcGbKGQApH5Aco4UAv+/q5nuie5V0qh7aHYl+wOqyA68zBF5vPlL5RDpBk8DzmWnhcHNbjhd7YcMc8U/8lwKUv/jpyPhw55WaaBA8cpebCmmIYGDrLiNQQmyLhZBobVTLfeYyGHqyd62EntzyyiEMS4tK+SKpkeI8K9qUza+KoGwVM/HdhfSHx96vW4salL7S67QUpWbub3ZlgnYbkffYnNH7KgFcHH1pvvSpdxzxjuxEapht2Q/OBLOL8SABBC8X2X7gBYMmE5Jan721cCGDPMFRmV+Rqsqi18Jmv21tF5ZSBTOOlBakw5jZNBuiBCR3ZwBW8eS9ochThM66am+LMTfGe7XPEDu1I80f1OWiCrTUFZ+OyudOgrw62bLJh4kV6HXBXuZQDi2cFGEychylZGdVOVdsNn2O+jgCdDMLMm3kliCiZMDFauoPE48plu3Xe9PVHYnQ1zijYN7lDE2Owdp1rVEXzvWx+pr5YpSVz1J/n7xTLYS0BS3ThYO0F8c5XHldsUgynuqZEbq+Nh7bzJLcoiGtY+HDxd68ar3Cv++sVY6X276KgftXp0kiLO2q1iBzWAHX4vzf84K/WFa2pH9UJI11VmS3iVZEHfARidw+E5hfnOek7aL71tvlR9jLGujl6ssBjjC5M6I/k5adKSVetLYW5KxsRlOJSuFMtDA2X5hcn258ucqkSSy23IM+fsDzjkSxQLY7y99a+S3QTGPxAm4WNY75A+bUUx9nc88ViZ4LLuehRqW79xrGkdJ3jwDAujmz0KxlHoYvgXVgG+1yAKRnEHlxvoFGSVaYSnE+nV/JgxKmAGXyH6z07U9qaLEDBN+RvD9JUiLoF8pGVJo9RNqI49mfBafVK/DR0iyjvqKfsYcfjvk7/rWtleVpCxm21PEMkcmFWng/jdEvx7AcIQO0V4rqP4CJL+uHvdqImEDsr3TzXHmThQzVgibjnLoa3D+b4OK2/D1/snzpEkVAd7hOqHgH3d7R5OHNxmPpt6ZbGYZ4ntQxEm0zp3pH/1pGbP6imP+llf6n1LsZYeCe4rsiss2pcRYWQF/LzfSSNOsHefggl7drEhnv5IMpdLF/JIPcClV41KYRjv1zXM3gXjj/79pjpO8ZGX9usQhi95tpait2+c6uZEGpUc6ysz3cDQadC+BfEqM7KKntpOAAAAD2gVwBfwL5U5+0jnWuwkRQL3G9swUgN1H1cLr84ReOB0MiEFYl8+JOOL1Tr8Xc3WLr6TuiUMx5/k+0czAjJDmG7SQ7LXGlygNmQ2FnSDNBY1kx3smn+sRnnHerK3Y+vwsznsK0oie6qhlYAhxNU1F+Vzshvyroe77kp05hI5SxzzBsE3jTRjrWNHvPDESToIpqIFyld/t4y7dL1iZ8ySwmo/E9NjJbDeWkvNwe3NlNQeQ2i3KH/dy2JgxYjIxupI+UdmXusMfJ5yzmYwEuOHzREcq0CCIo2SoYJYwCiq4wUhDZVnThw0UCUEgVs8MntIqBzrfcO1+HeuWaZ+2PiDK1dWbRwf4D9bY3hFaJS13CYRj32CUpYeB1mf4pK7cWN37DYzZMvtkEwDNhEsCFz03OedZVaLcAdkMtl0+/cDVTT54/82IFjxGKDzLcdZH1q4p4I9MN8qbjY/QL/72iYRYb6NDyialpvLUTEswus2AA50AwDpvHe7QLGnx75FvztcRdRaXhIAJtaww2d8y1oHsWjkcBcLvuDvzCWlGEGFe4cTBkqBTcDpbUkE88xjIVC7ChENMye7Ldm4ppl5NhIxX/XhMR6wYhOb/23xTeFSneYv+djNwjUo7YPmPkkwWR4nvexyadDVnb4RDdUvrzkimACsKSYGwnDvVCeS6KWYjSkhgn1hMFVXOlUB4ZbKuUTTTsSk3/SCOHpCRSK6/COAT2JP8WXrHJl89YrODCBZs0Y87BWmnD0jLg9FQi34qQn43emyjk0Tu8mAM/pBqsgNikxFUMKjcFWZgxXtZPuRd2V7b312h2T/6EuXCHfXIvhUY6+aeQmIqr0r7iaSo68AAAFHgq0depuEWXz9JeSfeyX8AoJJk4620w2uIDSn3nwr6bheqiRofUpHGfCCd+bgeV7JhHB/ztCfyA3mEp43CpbF5KK8GERM4fp96ARtiLYfuOMNUIKb417F8OQt7iFKMUFHXtzdSaF4AAGkWmXzBI5bk3OdUt1VvEtFf63C5MiQu5PQ47PWKH7wJezXgDp4dOi/7GdtMirWC9OklQX8J7r4njaM9FyD7RKyMzxNQWu/P3QOnWBaPcmmGn0DQCHqz8eKxPOknS8dP+D5OqU/MrR5+KJqwzGxAjvXyiTDPlMgl5i+qUTa7Pn2SNBS2P+RAFaSTymmBYndRb5LXJMI+EISxcEH2ik4+mnlpHDSI6JmLX6oSXjXKLWp6vyq2Aj/b/k4OMuumsxpk+GXME9kPjk36hUdGBpCZILxyGAVTpZUN6w1yIAvRPqRCsZp+pWyIdo1P8u7oUqkUDqF9soM6cXlwwMCG4RTXy2jDY3tRK/MqoLo+Adyglslu7J/RstJkMcJuMYI0EFiv9JlMal1f7Uc8mo9AwAKFbBks0lcjKR9ZKcxUZzz8wS6hAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAAA="
_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
def safe_styler_map(styler, func, subset=None):
    """سازگار با pandas 2.1+ و نسخه‌های قدیمی‌تر Styler."""
    if hasattr(styler, "map"):
        return styler.map(func, subset=subset)
    return styler.applymap(func, subset=subset)
def _text(value: Any) -> str:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    return str(value).translate(_FA_DIGITS).replace("\u200c", " ").replace("\xa0", " ").strip()
CATEGORY_MAP = {
    '3039': 'سس', '3061': 'سوسیس', '3062': 'کوکتل', '3064': 'وکیوم',
    '3065': 'هات داگ', '3066': 'رول گوشت', '3067': 'رول مرغ', '3068': 'پپرونی',
    '3070': 'رول های خاص', '3073': 'سوخاری', '3074': 'سیب زمینی', '3075': 'فلافل',
    '3077': 'همبرگر', '3078': 'کباب', '3081': 'پنیر مطهر', '3082': 'پنیر خمیرآوران',
    '3083': 'پنیر البان', '3084': 'چاشنی', '3087': 'ترشی', '3088': 'شوری',
    '3089': 'کنسرو گوشتی', '3090': 'کنسرو گیاهی', '3091': 'محصولات گوشتی',
    '3092': 'سبزی و غلات', '3093': 'رول مخلوط', '3095': 'ترشک', '3096': 'لواشک',
    '3097': 'چیپس', '3098': 'کالباس خانواده', '3100': 'رب میوه', '3103': 'پنیر سپاس',
    '3104': 'مارتا',
}
EXCLUDED_GROUPS = {"02", "09", "3071", "3080", "3099"}
EXCLUDED_BRANDS = {"250 گرمی", "250 گرمي", "300درصد گوشت", "300 درصد گوشت", "450-/+", "450+/-", "450-+/-", "650-/+", "650+/-", "99387258"}
def _is_excluded(value: Any) -> bool:
    """بررسی عضویت مقدار در برندها یا گروه‌های حذف‌شده."""
    t = _text(value)
    return t in EXCLUDED_BRANDS or t in EXCLUDED_GROUPS

def _is_excluded_group(value: Any) -> bool:
    """گروه‌های سیستمی/غیرکالایی را حتی با فاصله‌گذاری و ارقام متفاوت حذف می‌کند."""
    raw = _text(value).replace("ي", "ی").replace("ك", "ک")
    compact = _norm(raw)
    if not compact or compact in {"-", "—", "–"}:
        return True
    # گروه‌های درخواستی: -لاين فروش 0، -، 3099 -کارتخوان، 09 -اموال
    if "لاینفروش0" in compact or "لاينفروش0" in compact or "لاینفروش" in compact:
        return True
    if re.search(r"3099.*کارتخوان|کارتخوان.*3099", compact):
        return True
    if re.search(r"(?:^|[^0-9])0?9.*اموال|اموال.*(?:^|[^0-9])0?9", compact):
        return True
    return _code(raw) in EXCLUDED_GROUPS

def _is_valid_product_group(value: Any) -> bool:
    return not _is_excluded_group(value)
def _category_name(value: Any) -> str:
    """کد گروه را به نام فارسی قابل نمایش تبدیل می‌کند؛ مقدارهای ناشناخته حفظ می‌شوند."""
    label = _text(value)
    code = _code(value)
    if code in CATEGORY_MAP:
        return f"{CATEGORY_MAP[code]} ({code})"
    # پشتیبانی از کدهایی که همراه عنوان یا پرانتز در فایل آمده‌اند.
    match = re.search(r"(?<!\d)(\d{4})(?!\d)", label)
    if match and match.group(1) in CATEGORY_MAP:
        code = match.group(1)
        return f"{CATEGORY_MAP[code]} ({code})"
    return label
def _code(value: Any) -> str:
    """کد شناسه را رشتهٔ پاک نگه می‌دارد و پسوند عددی .0 را حذف می‌کند."""
    result = _text(value)
    if re.fullmatch(r"[+-]?\d+\.0+", result):
        result = result.split(".", 1)[0]
    return result.strip()
def _norm(value: Any) -> str:
    s = _text(value).lower().replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"[\s_\-/\\]+", "", s)
def _number(value: Any) -> float:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return np.nan
    if isinstance(value, (int, float, np.number)):
        return float(value)
    s = _text(value).replace(",", "").replace("،", "").replace("٬", "").replace("٫", ".").replace("%", "")
    try:
        return float(s)
    except (TypeError, ValueError):
        return np.nan
def _as_frame(data: Any) -> pd.DataFrame:
    """خواندن SpreadsheetML/XML، فایل xlsx/xlsm یا فایل باینری OLE2/xls."""
    if isinstance(data, pd.DataFrame):
        return data.copy()
    if hasattr(data, "getvalue"):
        data = data.getvalue()
    elif hasattr(data, "read"):
        stream = data
        pos = stream.tell() if hasattr(stream, "tell") else None
        data = stream.read()
        if pos is not None and hasattr(stream, "seek"):
            stream.seek(pos)
    raw = bytes(data)
    probe = raw.lstrip(b"\xef\xbb\xbf \t\r\n")
    if probe[:5].lower() == b"<?xml" or probe[:9].lower() == b"<workbook":
        return _spreadsheetml_to_frame(raw)
    try:
        if raw.startswith(b"PK\x03\x04"):
            engine = "openpyxl"
        elif raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
            engine = "xlrd"
        else:
            raise ValueError("هدر فایل نه SpreadsheetML/XML، نه xlsx و نه OLE2/xls است.")
        return pd.read_excel(io.BytesIO(raw), engine=engine, sheet_name=0, header=None, dtype=object)
    except Exception as exc:
        raise ValueError(f"فایل اکسل خوانده نشد: {exc}") from exc
def validate_uploaded_file(uploaded_file) -> tuple[bool, str]:
    """اعتبارسنجی فایل بارگذاری‌شده با توجه به پسوند و هدر واقعی؛ xls ممکن است SpreadsheetML باشد."""
    if uploaded_file is None:
        return False, "فایلی بارگذاری نشده است."
    name = getattr(uploaded_file, "name", "") or ""
    ext = Path(name).suffix.lower()
    raw = uploaded_file.getvalue() if hasattr(uploaded_file, "getvalue") else bytes(uploaded_file)
    probe = raw.lstrip(b"\xef\xbb\xbf \t\r\n")
    xml_header = probe[:5].lower() == b"<?xml" or probe[:9].lower() == b"<workbook"
    if ext not in {".xls", ".xlsx", ".xlsm", ".xml"}:
        return False, f"پسوند فایل پشتیبانی نمی‌شود: {ext or 'بدون پسوند'}"
    if xml_header:
        try:
            _spreadsheetml_to_frame(raw)
            return True, ""
        except Exception as exc:
            return False, f"فایل XML/SpreadsheetML معتبر نیست: {exc}"
    if ext == ".xml":
        return False, "فایل XML باید با <?xml یا <Workbook شروع شود."
    if raw.startswith(b"PK\x03\x04"):
        return (ext in {".xlsx", ".xlsm"}, "هدر ZIP فایل با پسوند اکسل سازگار نیست.") if ext not in {".xlsx", ".xlsm"} else (True, "")
    if raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
        return (ext == ".xls", "هدر OLE2 فایل با پسوند اکسل سازگار نیست.") if ext != ".xls" else (True, "")
    return False, "هدر فایل با قالب‌های SpreadsheetML/XML، xlsx/xlsm یا OLE2/xls سازگار نیست."
def _spreadsheetml_to_frame(raw: bytes) -> pd.DataFrame:
    """خواندن SpreadsheetML؛ ss:Index در ردیف و سلول هر دو 1-based است."""
    try:
        root = ET.fromstring(raw)
    except ET.ParseError as exc:
        raise ValueError(f"XML گزارش خرید معتبر نیست: {exc}") from exc
    ns_uri = "urn:schemas-microsoft-com:office:spreadsheet"
    ns = {"ss": ns_uri}
    ws = root.find(".//ss:Worksheet", ns)
    table = ws.find("ss:Table", ns) if ws is not None else None
    if table is None:
        raise ValueError("جدول Worksheet در فایل XML پیدا نشد.")
    rows: list[list[Any]] = []
    next_row = 1
    for row_el in table.findall("ss:Row", ns):
        row_idx = int(row_el.attrib.get(f"{{{ns_uri}}}Index", next_row))
        while len(rows) < row_idx - 1:
            rows.append([])
        row: list[Any] = []
        next_col = 1
        for cell in row_el.findall("ss:Cell", ns):
            col_idx = int(cell.attrib.get(f"{{{ns_uri}}}Index", next_col))
            while len(row) < col_idx - 1:
                row.append(None)
            data_el = cell.find("ss:Data", ns)
            val: Any = data_el.text if data_el is not None else None
            dtype = data_el.attrib.get(f"{{{ns_uri}}}Type", "") if data_el is not None else ""
            if val is not None and dtype == "Number":
                try: val = float(val)
                except (ValueError, TypeError): pass
            row.append(val)
            next_col = col_idx + 1
        rows.append(row)
        next_row = row_idx + 1
    width = max((len(r) for r in rows), default=0)
    return pd.DataFrame([r + [None] * (width-len(r)) for r in rows])
def _find_header(frame: pd.DataFrame, required: tuple[str, ...], scan: int = 80) -> int:
    # سلول‌ها پیش از تطبیق normalize می‌شوند؛ توکن‌های لازم نیز باید همان‌طور
    # normalize شوند، وگرنه فاصلهٔ «کد کالا» با مقدار «کدکالا» تطبیق نمی‌کند.
    tokens = tuple(_norm(token) for token in required)
    best = (-1, -1)
    for i in range(min(len(frame), scan)):
        vals = [_norm(v) for v in frame.iloc[i].tolist()]
        score = sum(any(token in v for token in tokens) for v in vals if v)
        if score > best[1]:
            best = (i, score)
        if score >= min(2, len(tokens)):
            return i
    if best[1] > 0:
        return best[0]
    raise ValueError("سطر سرستون قابل شناسایی نیست.")
def _columns_from_header(frame: pd.DataFrame, row: int) -> pd.DataFrame:
    headers = []
    for j, v in enumerate(frame.iloc[row].tolist()):
        h = _text(v) or f"ستون_{j+1}"
        headers.append(h)
    # نام ستون تکراری مانع انتخاب/نگاشت در pandas نشود.
    seen: dict[str, int] = {}
    unique = []
    for h in headers:
        seen[h] = seen.get(h, 0) + 1
        unique.append(h if seen[h] == 1 else f"{h}_{seen[h]}")
    out = frame.iloc[row+1:].copy()
    out.columns = unique
    return out.reset_index(drop=True)
def _col(columns, *terms):
    pairs = [(_norm(c), c) for c in columns]
    for term in terms:
        t = _norm(term)
        for n, c in pairs:
            if n == t:
                return c
    for term in terms:
        t = _norm(term)
        for n, c in pairs:
            if t and t in n:
                return c
    return None
def _date_key(value: Any):
    """کلید مرتب‌سازی تاریخ میلادی یا شمسی متنی؛ برای تاریخ نامشخص None."""
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return (value.year, value.month, value.day)
    s = _text(value)
    m = re.search(r"(\d{4})\s*[/\-.年]\s*(\d{1,2})\s*[/\-.月]\s*(\d{1,2})", s)
    if m:
        return tuple(map(int, m.groups()))
    try:
        dt = pd.to_datetime(s, errors="coerce")
        if pd.notna(dt):
            return (dt.year, dt.month, dt.day)
    except Exception:
        pass
    return None
def parse_purchase_report(file_bytes_or_buffer) -> pd.DataFrame:
    """آخرین خرید هر کد کالا را برمی‌گرداند؛ فیلتر نوع فاکتور خرید و محاسبه بهای تمام‌شده.

    ستون‌های خروجی: کد کالا، آخرین_قیمت_خرید، آخرین_تخفیف، قیمت_تمام_شده، تاریخ_خرید، طرف_حساب.
    آخرین تاریخ معتبر اولویت دارد؛ در تساوی یا نبود تاریخ، آخرین ردیف فایل انتخاب می‌شود.
    """
    raw_frame = _as_frame(file_bytes_or_buffer)
    hdr = _find_header(raw_frame, ("کد کالا", "قیمت واحد"))
    frame = _columns_from_header(raw_frame, hdr)
    code = _col(frame.columns, "کد کالا")
    # در بعضی گزارش‌های خرید، به‌علت شیفت/ادغام سلول‌ها، مقدار کد کالا
    # در همان اندیس سرستون یا یک ستون پس از آن (col_idx یا col_idx+1) قرار
    # دارد. هر دو ستون نامزد بررسی می‌شوند و ستونی که کدهای معتبر بیشتری
    # دارد انتخاب می‌شود.
    header_positions = [
        j for j, value in enumerate(raw_frame.iloc[hdr].tolist())
        if _norm(value) == _norm("کد کالا")
    ]
    code_candidates: list = []
    for j in header_positions:
        for candidate_idx in (j, j + 1):
            if candidate_idx < len(frame.columns):
                candidate = frame.columns[candidate_idx]
                if candidate not in code_candidates:
                    code_candidates.append(candidate)
    if code_candidates:
        def _code_count(column) -> int:
            return int(frame[column].map(_code).ne("").sum())
        best_candidate = max(code_candidates, key=_code_count)
        best_count = _code_count(best_candidate)
        current_count = _code_count(code) if code is not None else -1
        if code is None or best_count > current_count:
            code = best_candidate
    unit = _col(frame.columns, "قیمت واحد")
    discount = _col(frame.columns, "تخفیف واحد")
    invoice_type = _col(frame.columns, "نوع فاکتور")
    dcol = _col(frame.columns, "تاریخ")
    party = _col(frame.columns, "نام طرف حساب", "طرف حساب")
    if not code or not unit:
        raise ValueError("ستون‌های «کد کالا» و «قیمت واحد» در گزارش خرید پیدا نشد.")
    if invoice_type is None:
        raise ValueError("ستون «نوع فاکتور» در گزارش خرید پیدا نشد.")
    keep = frame[invoice_type].map(lambda x: _norm(x) in {"خرید", "خريد"})
    f = frame.loc[keep].copy()
    f["کد کالا"] = f[code].map(_code)
    f["آخرین_قیمت_خرید"] = f[unit].map(_number)
    f["آخرین_تخفیف"] = f[discount].map(_number) if discount else 0.0
    f["آخرین_تخفیف"] = f["آخرین_تخفیف"].fillna(0)
    f["قیمت_تمام_شده"] = f["آخرین_قیمت_خرید"] - f["آخرین_تخفیف"]
    f["تاریخ_خرید"] = f[dcol] if dcol else None
    f["طرف_حساب"] = f[party] if party else None
    f["_source_order"] = np.arange(len(frame))[keep.to_numpy()]
    f["_date_key"] = f["تاریخ_خرید"].map(_date_key)
    # در هر کد، تاریخ‌های معتبر بر تاریخ نامعتبر مقدم‌اند و ترتیب فایل tie-breaker است.
    f["_sort_date"] = f["_date_key"].map(lambda x: x if x is not None else (-1,-1,-1))
    f = f.sort_values(["_sort_date", "_source_order"], kind="stable")
    latest = f.drop_duplicates("کد کالا", keep="last")
    return latest[["کد کالا", "آخرین_قیمت_خرید", "آخرین_تخفیف", "قیمت_تمام_شده", "تاریخ_خرید", "طرف_حساب"]].reset_index(drop=True)
PRODUCT_COLUMNS = ["کد کالا", "عنوان کالا", "برند", "عنوان گروه", "قیمت فروش", "قیمت 2", "قیمت 3", "قیمت 4", "قیمت 5", "ق مصرف کننده"]
def parse_product_list(file_bytes_or_buffer, header_row: int | None = None):
    """خواندن لیست کالا با کشف خودکار سربرگ (ردیف 4 پیش‌فرض)، حذف غیرفعال‌ها.
    خروجی (DataFrame استاندارد، metadata) شامل شماره دقیق سطر اکسل برای خروجی است.
    """
    raw_frame = _as_frame(file_bytes_or_buffer)
    if header_row is None:
        try:
            hdr = _find_header(raw_frame, ("کد کالا", "عنوان کالا"))
        except ValueError:
            hdr = min(3, max(0, len(raw_frame)-1))
    else:
        hdr = max(0, int(header_row)-1)
    frame = _columns_from_header(raw_frame, hdr)
    cols = frame.columns
    aliases = {
        "کد کالا": ("کد کالا/خدمات", "کد کالا", "کد"),
        "عنوان کالا": ("عنوان کالا/خدمات", "عنوان کالا", "نام کالا"),
        "قیمت فروش": ("قیمت فروش",), "قیمت 2": ("قیمت 2",),
        "قیمت 3": ("قیمت 3",), "قیمت 4": ("قیمت 4",), "قیمت 5": ("قیمت 5",),
        "ق مصرف کننده": ("ق مصرف کننده", "قیمت مصرف کننده", "مصرف کننده"),
        "غیر فعال": ("غیر فعال", "غیرفعال"), "عنوان گروه": ("عنوان گروه", "گروه"),
    }
    source = {std: _col(cols, *names) for std, names in aliases.items()}
    if source["کد کالا"] is None or source["عنوان کالا"] is None:
        raise ValueError("ستون کد کالا یا عنوان کالا در لیست پیدا نشد.")
    result = pd.DataFrame(index=frame.index)
    for std in PRODUCT_COLUMNS:
        result[std] = frame[source[std]] if source.get(std) else (np.nan if std not in ("کد کالا", "عنوان کالا", "برند") else "")
    result["کد کالا"] = result["کد کالا"].map(_code)
    result["عنوان کالا"] = result["عنوان کالا"].map(_text)
    result["عنوان گروه"] = result["عنوان گروه"].map(_category_name)
    result["برند"] = result["عنوان کالا"].str.extract(r"\(([^()]*)\)\s*$", expand=False).fillna("").str.strip()
    inactive_col = source.get("غیر فعال")
    def is_true(x):
        return x is True or _norm(x) in {"true", "1", "بله", "yes", "y"}
    inactive = frame[inactive_col].map(is_true) if inactive_col else pd.Series(False, index=frame.index)
    result["_excel_row"] = np.arange(len(frame), dtype=int) + hdr + 2
    result = result.loc[~inactive & result["کد کالا"].ne("")].reset_index(drop=True)
    metadata = {"header_row": hdr+1, "source_columns": source, "original_columns": list(cols)}
    return result, metadata
def markup(buy, sell):
    b, s = _number(buy), _number(sell)
    return ((s-b)/b*100) if np.isfinite(b) and b > 0 and np.isfinite(s) else np.nan
def _money(x):
    try:
        return f"{float(x):,.0f}"
    except (TypeError, ValueError):
        return "—"
def _style_profit(value, target=30.0, tolerance=5.0):
    """سبز تیره بالاتر از بازه مطلوب، سبز روشن داخل تلورانس و قرمز پایین‌تر."""
    try:
        number, target, tolerance = float(value), float(target), max(0.0, float(tolerance))
    except (TypeError, ValueError):
        number, target, tolerance = np.nan, 30.0, 5.0
    if not np.isfinite(number):
        return "color:#777777"
    if number < target - tolerance:
        color = "#c62828"
    elif number <= target + tolerance:
        color = "#8bcf9b"
    else:
        color = "#188038"
    return f"color:{color};font-weight:bold"

def render_color_guide(target=30.0, tolerance=5.0):
    """راهنمای رنگ متناسب با حد مطلوب و تلورانس جاری."""
    low, high = max(0.0, target - tolerance), target + tolerance
    with st.expander("🎨 راهنمای رنگ‌ها", expanded=False):
        st.markdown(f"""<div class="color-guide" style="display:flex;gap:16px;flex-wrap:wrap;align-items:center">
          <span><i style="display:inline-block;width:12px;height:12px;border-radius:50%;background:#188038"></i> سبز تیره: بیشتر از {high:g}٪</span>
          <span><i style="display:inline-block;width:12px;height:12px;border-radius:50%;background:#8bcf9b"></i> سبز روشن: {low:g}٪ تا {high:g}٪</span>
          <span><i style="display:inline-block;width:12px;height:12px;border-radius:50%;background:#c62828"></i> قرمز: کمتر از {low:g}٪</span>
          <span><i style="display:inline-block;width:12px;height:12px;border-radius:50%;background:#7B2CBF"></i> بنفش: پیشنهاد قیمت مصرف‌کننده</span>
          <span><i style="display:inline-block;width:12px;height:12px;border-radius:50%;background:#777777"></i> خاکستری: ناموجود/نامعتبر</span>
        </div>""", unsafe_allow_html=True)

def _to_xlsx(original: bytes, filename: str, sheet_name: str, edits: pd.DataFrame, meta: dict) -> bytes:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "xls":
        raise ValueError("خروجی ویرایشی xls ممکن نیست؛ فایل فهرست را به xlsx تبدیل و دوباره بارگذاری کنید.")
    wb = openpyxl.load_workbook(io.BytesIO(original), keep_vba=(ext == "xlsm"), data_only=False)
    ws = wb[sheet_name]
    src = meta["source_columns"]
    # Build the header-to-column map once; avoid rescanning every header per cell.
    headers = [_text(ws.cell(meta["header_row"], c).value) for c in range(1, ws.max_column + 1)]
    header_to_col = {}
    for col_idx, header in enumerate(headers, start=1):
        if header and header not in header_to_col:
            header_to_col[header] = col_idx
    export_cols = {**src}
    for idx, row in edits.iterrows():
        excel_row = int(row["_excel_row"])
        for standard in PRODUCT_COLUMNS:
            if standard in ("کد کالا", "عنوان کالا") or standard not in edits.columns:
                continue
            original_header = export_cols.get(standard)
            if not original_header:
                continue
            # عنوان سربرگ منحصربه‌فرد به شماره ستون؛ ستون‌های خالی/تکراری هم پشتیبانی می‌شوند.
            col_idx = header_to_col.get(_text(original_header))
            if col_idx is None:
                continue
            value = row[standard]
            if pd.isna(value):
                value = None
            elif isinstance(value, str):
                value = _number(value)
                if not np.isfinite(value):
                    value = row[standard]
            ws.cell(excel_row, col_idx).value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
def apply_markup_columns(frame: pd.DataFrame) -> pd.DataFrame:
    """محاسبهٔ مجدد ستون‌های مارک‌آپ پس از تغییر قیمت‌ها."""
    buy = pd.to_numeric(frame["قیمت تمام شده خرید"], errors="coerce")
    def _vector_markup(sell_col):
        sell = pd.to_numeric(frame[sell_col], errors="coerce")
        return ((sell - buy) / buy * 100).where(buy.gt(0) & sell.notna())
    frame["مارک‌آپ فروش ٪"] = _vector_markup("قیمت فروش")
    frame["مارک‌آپ عمده ٪"] = _vector_markup("قیمت 3")
    frame["مارک‌آپ مصرف‌کننده ٪"] = _vector_markup("ق مصرف کننده").where(pd.to_numeric(frame["ق مصرف کننده"], errors="coerce").gt(0))
    return frame


def show_pricing_page():
    """Render the pricing and Excel editing page."""
    st.markdown(f"""<div class="hero"><h1><img src="data:image/webp;base64,{LOGO_B64}" alt="لوگوی پروتو بازار" style="width:54px;height:54px;object-fit:contain;vertical-align:middle;margin-left:10px"> ماژول قیمت‌گذاری پروتو بازار</h1><p>تحلیل قیمت، بررسی حاشیه سود و ویرایش امن فایل اکسل</p></div>""", unsafe_allow_html=True)
    st.caption("پردازش خودکار فهرست کالا و سابقهٔ خرید، تحلیل حاشیهٔ فروش و خروجی اکسلِ قابل ویرایش")
    for key, default in {"processed":False,"products":None,"purchase":None,"meta":None,"raw":None,"raw_name":None,"sheet":None,"edited":None,"source_signature":None}.items():
        if key not in st.session_state: st.session_state[key] = default
    with st.sidebar:
        st.markdown(f'<div class="sidebar-brand-card"><img src="data:image/webp;base64,{LOGO_B64}" alt="Protobazar"><div><strong>پروتو بازار</strong><small>سامانه هوشمند قیمت‌گذاری</small></div></div>', unsafe_allow_html=True)
        st.subheader("بارگذاری فایل‌ها")
        products_up = st.file_uploader("۱. فهرست کالاها (xlsx / xlsm)", type=["xlsx","xlsm","xls"], key="products_file", help="مسیر استخراج از کیان‌پرداز: اطلاعات پایه > انبار/کالا > معرفی کالا")
        purchase_up = st.file_uploader("۲. گزارش خرید (xls / xlsx / XML)", type=["xls","xlsx","xml"], key="purchase_file", help="مسیر استخراج از کیان‌پرداز: گزارش > گزارش خرید و فروش کالا > آمار خرید کالا (ریز عملکرد)")
        st.caption("گزارش‌های xls استاندارد و SpreadsheetML/XML هر دو خوانده می‌شوند.")
    prod_bytes = products_up.getvalue() if products_up else None
    purchase_bytes = purchase_up.getvalue() if purchase_up else None
    products_valid, products_validation_message = validate_uploaded_file(products_up) if products_up else (False, "")
    purchase_valid, purchase_validation_message = validate_uploaded_file(purchase_up) if purchase_up else (False, "")
    if products_up and not products_valid:
        st.sidebar.error(f"فهرست کالا: {products_validation_message}")
    if purchase_up and not purchase_valid:
        st.sidebar.error(f"گزارش خرید: {purchase_validation_message}")
    files_valid = bool(prod_bytes and purchase_bytes and products_valid and purchase_valid)
    signature = (products_up.name if products_up else None, (__import__("hashlib").sha256(prod_bytes).hexdigest() if prod_bytes else None),
                 purchase_up.name if purchase_up else None, (__import__("hashlib").sha256(purchase_bytes).hexdigest() if purchase_bytes else None))
    process_now = st.button("🚀 شروع پردازش و تحلیل قیمت‌ها", type="primary", use_container_width=True,
                            disabled=not files_valid)
    auto_process = bool(files_valid and signature != st.session_state.source_signature)
    if process_now or auto_process:
        try:
            @st.cache_data(show_spinner=False)
            def _cached_parse(prod_b, purch_b):
                return parse_product_list(prod_b), parse_purchase_report(purch_b)
            (p_df, p_meta), h_df = _cached_parse(prod_bytes, purchase_bytes)
            p_df["_key"] = p_df["کد کالا"].map(_code).map(_norm)
            h_df["_key"] = h_df["کد کالا"].map(_code).map(_norm)
            costs = h_df.drop_duplicates("_key", keep="last").set_index("_key")
            p_df["قیمت تمام شده خرید"] = p_df["_key"].map(costs["قیمت_تمام_شده"])
            p_df["آخرین تخفیف خرید"] = p_df["_key"].map(costs["آخرین_تخفیف"])
            p_df["تاریخ خرید"] = p_df["_key"].map(costs["تاریخ_خرید"])
            p_df["طرف حساب"] = p_df["_key"].map(costs["طرف_حساب"])
            for price in ("قیمت فروش","قیمت 2","قیمت 3","قیمت 4","قیمت 5","ق مصرف کننده"):
                p_df[price] = p_df[price].map(_number)
            p_df = apply_markup_columns(p_df)
            p_df["قیمت پیشنهادی مصرف‌کننده"] = np.nan
            buy_ok = p_df["قیمت تمام شده خرید"].notna() & (p_df["قیمت تمام شده خرید"] > 0)
            consumer_ok = p_df["ق مصرف کننده"].notna() & (p_df["ق مصرف کننده"] > 0) & buy_ok
            ratios = p_df.loc[consumer_ok].assign(_ratio=lambda x:x["ق مصرف کننده"]/x["قیمت تمام شده خرید"])
            global_ratio = ratios["_ratio"].median() if len(ratios) else 1.35
            missing = p_df["ق مصرف کننده"].isna() | (p_df["ق مصرف کننده"] <= 0)
            for idx in p_df.index[missing & buy_ok]:
                brand = p_df.at[idx,"برند"]
                br = ratios.loc[ratios["برند"].eq(brand),"_ratio"] if brand else pd.Series(dtype=float)
                ratio = float(br.median()) if len(br) else float(global_ratio)
                p_df.at[idx,"قیمت پیشنهادی مصرف‌کننده"] = round(float(p_df.at[idx,"قیمت تمام شده خرید"])*ratio)
            p_df.drop(columns="_key", inplace=True)
            st.session_state.products, st.session_state.purchase = p_df, h_df
            st.session_state.meta, st.session_state.raw = p_meta, prod_bytes
            st.session_state.raw_name = products_up.name
            st.session_state.sheet = "Sheet1"
            # شناسایی نام شیت از فایل اصلی
            if products_up.name.lower().endswith(("xlsx","xlsm")):
                book = openpyxl.load_workbook(io.BytesIO(prod_bytes), read_only=True, keep_vba=products_up.name.lower().endswith("xlsm"))
                st.session_state.sheet = book.sheetnames[0]
            st.session_state.edited = p_df.copy()
            st.session_state.source_signature = signature
            st.session_state.processed = True
            matched = int(p_df["قیمت تمام شده خرید"].notna().sum())
            st.success(f"پردازش موفق: {len(p_df):,} کالای فعال، {len(h_df):,} کد خرید یکتا؛ {matched:,} کالا با بهای خرید تطبیق یافت.")
        except Exception as exc:
            st.session_state.processed = False
            st.error(f"پردازش فایل‌ها ناموفق بود: {exc}")
    if not prod_bytes or not purchase_bytes:
        st.info("هر دو فایل را بارگذاری کنید. با بارگذاری فایل دوم پردازش خودکار انجام می‌شود؛ یا دکمهٔ شروع را بزنید.")
    elif not st.session_state.processed:
        st.warning("فایل‌ها حاضرند؛ برای پردازش دکمهٔ شروع را بزنید.")
    if st.session_state.processed and st.session_state.products is not None:
        df = st.session_state.products.copy()
        # گروه‌های سیستمی از کل نماهای ماژول حذف می‌شوند، اما ردیف‌های فایل اصلی برای حفظ ساختار خروجی دست‌نخورده می‌مانند.
        df = df.loc[~df["عنوان گروه"].map(_is_excluded_group)].copy()
        tabs = st.tabs(["📊 تحلیل", "✏️ ویرایش قیمت", "💡 پیشنهاد قیمت مصرف‌کننده", "💾 خروجی اکسل"])
        with tabs[0]:
            a_search, a_brand, a_group = st.columns([2,1,1])
            query = a_search.text_input("🔍 جستجو در عنوان کالا", key="analysis_search").strip()
            brands = sorted([x for x in df["برند"].dropna().unique() if x and not _is_excluded(x)])
            groups = sorted({str(x).strip() for x in df["عنوان گروه"].dropna() if _text(x) and _is_valid_product_group(x)})
            selected = a_brand.multiselect("فیلتر برند", brands, key="analysis_brand")
            selected_groups = a_group.multiselect("فیلتر دسته‌بندی (عنوان گروه)", groups, key="analysis_group")
            # حد مطلوب کلی و امکان تعیین حد جداگانه برای هر برند
            c_target, c_brand_target = st.columns([1, 2])
            default_target = c_target.number_input("حد مارک‌آپ مطلوب (٪)", min_value=0.0, max_value=500.0, value=30.0, step=1.0, key="markup_target_default")
            with c_brand_target:
                with st.expander("تنظیم حد مطلوب برای برندهای خاص"):
                    brand_targets = st.session_state.get("markup_brand_targets", {})
                    target_brand = st.selectbox("برند", ["(حد پیش‌فرض)"] + brands, key="markup_target_brand")
                    if target_brand != "(حد پیش‌فرض)":
                        brand_value = st.number_input("حد مطلوب این برند (٪)", min_value=0.0, max_value=500.0, value=float(brand_targets.get(target_brand, default_target)), step=1.0, key="markup_target_brand_value")
                        if st.button("ثبت حد برای برند", key="save_brand_markup_target"):
                            brand_targets[target_brand] = brand_value
                            st.session_state.markup_brand_targets = brand_targets
                            st.rerun()
                    if brand_targets:
                        st.caption("حدهای اختصاصی: " + "، ".join(f"{k}: {v:g}٪" for k,v in brand_targets.items()))
            tolerance = st.number_input("تلورانس حد مطلوب (± درصد)", min_value=0.0, max_value=100.0, value=5.0, step=1.0, key="markup_tolerance")
            # همان view برای KPIها، جدول و نمودار استفاده می‌شود تا همهٔ آمار با فیلترها سازگار باشد.
            view = df.copy()
            if selected:
                view = view[view["برند"].isin(selected)]
            if selected_groups:
                view = view[view["عنوان گروه"].map(lambda x: str(x).strip()).isin(selected_groups)]
            if query:
                view = view[view["عنوان کالا"].str.contains(query, case=False, na=False)]
            targets = st.session_state.get("markup_brand_targets", {})
            target_for_row = view["برند"].map(lambda b: float(targets.get(b, default_target)))
            view["حد مطلوب مارک‌آپ ٪"] = target_for_row
            view["وضعیت هشدار"] = ["نامعتبر" if not np.isfinite(pd.to_numeric(m, errors="coerce")) else ("هشدار" if float(m) < max(0.0, float(t)-tolerance) else ("مطلوب - بازه تلورانس" if float(m) <= float(t)+tolerance else "مطلوب")) for m,t in zip(view["مارک‌آپ فروش ٪"], target_for_row)]
            alert_filter = st.selectbox("نمایش وضعیت", ["همه کالاها", "فقط هشدارها", "هشدار و نیازمند بررسی", "فقط مطلوب"], key="analysis_alert_filter")
            if alert_filter == "فقط هشدارها":
                view = view[view["وضعیت هشدار"].eq("هشدار")]
            elif alert_filter == "هشدار و نیازمند بررسی":
                view = view[view["وضعیت هشدار"].isin(["هشدار", "نامعتبر"])]
            elif alert_filter == "فقط مطلوب":
                view = view[view["وضعیت هشدار"].isin(["مطلوب", "مطلوب - بازه تلورانس"])]
            cost = view["قیمت تمام شده خرید"]
            markup_values = pd.to_numeric(view["مارک‌آپ فروش ٪"], errors="coerce").to_numpy(dtype=float)
            valid_markup = markup_values[np.isfinite(markup_values) & (markup_values >= 0) & (markup_values <= 500)]
            a,b,c,d = st.columns(4)
            a.metric("کالاهای فعال",f"{len(view):,}")
            b.metric("دارای بهای خرید",f"{cost.notna().sum():,}")
            c.metric("میانگین مارک‌آپ فروش (۰ تا ۵۰۰٪)",f"{valid_markup.mean():.1f}٪" if len(valid_markup) else "—")
            d.metric("فاقد قیمت مصرف‌کننده",f"{(view['ق مصرف کننده'].isna() | (view['ق مصرف کننده']<=0)).sum():,}")
            st.caption(f"نمایش {len(view):,} کالا از {len(df):,}؛ میانگین مارک‌آپ فقط از مقادیر متناهی ۰ تا ۵۰۰٪ محاسبه می‌شود.")
            render_color_guide(default_target, tolerance)
            view = view.copy()
            no_consumer = pd.to_numeric(view["ق مصرف کننده"], errors="coerce").isna() | pd.to_numeric(view["ق مصرف کننده"], errors="coerce").le(0)
            view.loc[no_consumer, "مارک‌آپ مصرف‌کننده ٪"] = np.nan
            view.loc[no_consumer, "مارک‌آپ مصرف‌کننده ٪"] = "نیازمند قیمت مصرف"
            display_cols = ["کد کالا","عنوان کالا","برند","عنوان گروه","تاریخ خرید","قیمت تمام شده خرید","قیمت فروش","قیمت 3","ق مصرف کننده","قیمت پیشنهادی مصرف‌کننده","مارک‌آپ فروش ٪","مارک‌آپ عمده ٪","مارک‌آپ مصرف‌کننده ٪","حد مطلوب مارک‌آپ ٪","وضعیت هشدار"]
            def _row_markup_style(value):
                return _style_profit(value, default_target, tolerance)
            st.dataframe(safe_styler_map(view[display_cols].style, _row_markup_style, subset=["مارک‌آپ فروش ٪","مارک‌آپ عمده ٪"]).format({c:"{:,.0f}" for c in ["قیمت تمام شده خرید","قیمت فروش","قیمت 3","ق مصرف کننده","قیمت پیشنهادی مصرف‌کننده"]},na_rep="—"),use_container_width=True,height=450,hide_index=True)
            if len(valid_markup):
                st.subheader("آمار مارک‌آپ قیمت فروش (داده‌های فیلترشده)")
                chart_data = (pd.Series(valid_markup).round(1).value_counts().sort_index()
                              .rename_axis("درصد مارک‌آپ").rename("تعداد کالا").to_frame())
                st.bar_chart(chart_data, use_container_width=True)
            else:
                st.info("برای فیلتر فعلی، مارک‌آپ معتبر در بازهٔ ۰ تا ۵۰۰٪ وجود ندارد.")
        with tabs[1]:
            render_color_guide()
            st.markdown("قیمت‌های ویرایش‌شده تنها هنگام دانلود، در سلول‌های متناظر فایل اصلی نوشته می‌شوند.")
            price_edit=["قیمت فروش","قیمت 2","قیمت 3","قیمت 4","قیمت 5","ق مصرف کننده"]
            edited_state = st.session_state.edited
            brand_list = sorted([x for x in edited_state["برند"].dropna().unique() if x and not _is_excluded(x)])
            group_list = sorted({str(x).strip() for x in edited_state["عنوان گروه"].dropna() if _text(x) and _is_valid_product_group(x)})
            # --- جستجو و فیلتر ---
            f_search, f_brand, f_group = st.columns([2,1,1])
            edit_query = f_search.text_input("🔍 جستجو در عنوان کالا", key="edit_search").strip()
            edit_brands = f_brand.multiselect("فیلتر برند", brand_list, key="edit_brand_filter")
            edit_groups = f_group.multiselect("فیلتر دسته‌بندی (عنوان گروه)", group_list, key="edit_group_filter")
            editable=edited_state[["_excel_row","کد کالا","عنوان کالا","برند","عنوان گروه"]+price_edit].copy()
            editable["عنوان گروه"]=editable["عنوان گروه"].map(_text)
            editable = editable.loc[~editable["عنوان گروه"].map(_is_excluded_group)].copy()
            if edit_brands:
                editable = editable[editable["برند"].isin(edit_brands)]
            if edit_groups:
                editable = editable[editable["عنوان گروه"].isin(edit_groups)]
            if edit_query:
                editable = editable[editable["عنوان کالا"].str.contains(edit_query, case=False, na=False)]
            st.caption(f"نمایش {len(editable):,} کالا از {len(edited_state):,}")
            edited_view=st.data_editor(editable,hide_index=True,num_rows="fixed",use_container_width=True,height=420,
                                       disabled=["_excel_row","کد کالا","عنوان کالا","برند","عنوان گروه"],key="price_editor")
            # هر بار اجرا مقدار داده‌ویرایشگر به وضعیت پایدار نگاشت شود.
            for col in price_edit:
                st.session_state.edited.loc[edited_view.index,col]=pd.to_numeric(edited_view[col],errors="coerce")

            # --- ابزار تغییر گروهی قیمت (Bulk Price Update) ---
            st.subheader("⚙️ ابزار تغییر گروهی قیمت")
            st.caption("محدوده را انتخاب کنید؛ سپس درصد تغییر و ستون قیمت را تعیین کنید. فیلتر برند/دسته فقط در محدودهٔ مربوط فعال می‌شود.")
            b1,b2,b3 = st.columns(3)
            scope = b1.selectbox("محدوده تغییر", ["همۀ کالاها", "بر اساس برند", "بر اساس دسته‌بندی"], key="bulk_scope")
            bulk_brand = b2.selectbox("برند (مثلاً هایزم)", brand_list, disabled=(scope!="بر اساس برند"), key="bulk_brand")
            bulk_group = b3.selectbox("دسته‌بندی", group_list, disabled=(scope!="بر اساس دسته‌بندی"), key="bulk_group")
            b4,b5,b6 = st.columns(3)
            bulk_col = b4.selectbox("ستون قیمت هدف", price_edit, key="bulk_col")
            bulk_percent = b5.number_input("درصد تغییر (+10 یعنی افزایش ۱۰٪ و -5 یعنی کاهش ۵٪)", value=0.0, step=0.5, key="bulk_percent")
            bulk_round = b6.selectbox("گرد کردن قیمت نهایی", ["بدون گرد کردن","گرد به 500","گرد به 1000","گرد به 10000"], key="bulk_round")
            base_options = ["قیمت فروش","قیمت 2","قیمت 3","قیمت 4","قیمت 5","قیمت تمام شده خرید","ق مصرف کننده"]
            bulk_base = st.selectbox("قیمت پایه (مبنای محاسبهٔ درصد تغییر)", base_options, key="bulk_base",
                                     help="مثال: +10 یعنی ۱۰٪ افزایش نسبت به «قیمت پایه» انتخابی؛ -5 یعنی ۵٪ کاهش نسبت به همان ستون.")
            apply_bulk = st.button("✅ اعمال تغییرات درصدی به قیمت‌ها", type="primary", use_container_width=True)
            if apply_bulk:
                target_state = st.session_state.edited
                mask = pd.Series(True, index=target_state.index)
                if scope=="بر اساس برند":
                    mask &= target_state["برند"]==bulk_brand
                elif scope=="بر اساس دسته‌بندی":
                    mask &= target_state["عنوان گروه"].map(_text)==bulk_group
                if bulk_percent==0:
                    st.warning("درصد تغییر صفر است؛ تغییری اعمال نشد.")
                elif int(mask.sum())==0:
                    st.warning("هیچ کالایی با این محدوده پیدا نشد.")
                else:
                    target_vals = pd.to_numeric(target_state.loc[mask, bulk_col], errors="coerce")
                    if bulk_base == bulk_col:
                        base_vals = target_vals
                    else:
                        base_vals = pd.to_numeric(target_state.loc[mask, bulk_base], errors="coerce")
                    changed = base_vals.notna() & (base_vals>0)
                    new_vals = base_vals[changed] * (1 + bulk_percent/100.0)
                    if bulk_round.startswith("گرد"):
                        step = int(bulk_round.split()[-1])
                        new_vals = (new_vals/step).round()*step
                    # «بدون گرد کردن» یعنی حفظ مقدار محاسبه‌شدهٔ دقیق.
                    st.session_state.edited.loc[new_vals.index, bulk_col] = new_vals
                    st.session_state.edited = apply_markup_columns(st.session_state.edited)
                    st.success(f"✅ قیمت «{bulk_col}» بر مبنای «{bulk_base}» برای {len(new_vals):,} کالا با {bulk_percent:+.1f}٪ تغییر کرد"
                               + (f" و به مضرب {bulk_round.split()[-1]} گرد شد." if bulk_round.startswith('گرد') else "."))
                    st.rerun()
            if st.button("بازنشانی ویرایش‌ها به فایل پردازش‌شده"):
                st.session_state.edited=st.session_state.products.copy()
                st.rerun()
        with tabs[2]:
            st.subheader("پیشنهاد قیمت مصرف‌کننده برای کالاهای فاقد قیمت")
            missing = df["ق مصرف کننده"].isna() | (df["ق مصرف کننده"]<=0)
            proposals=df.loc[missing & df["قیمت پیشنهادی مصرف‌کننده"].notna(),["کد کالا","عنوان کالا","برند","قیمت تمام شده خرید","قیمت پیشنهادی مصرف‌کننده"]].copy()
            st.caption("قیمت مصرف‌کنندهٔ موجود دست‌نخورده می‌ماند. پیشنهاد بنفش با نسبت میانهٔ همان برند ساخته می‌شود؛ در نبود نمونهٔ برند، میانهٔ کل استفاده می‌شود.")
            st.dataframe(proposals.style.set_properties(subset=["قیمت پیشنهادی مصرف‌کننده"],**{"color":"#7B2CBF","font-weight":"bold"}).format({"قیمت تمام شده خرید":"{:,.0f}","قیمت پیشنهادی مصرف‌کننده":"{:,.0f}"},na_rep="—"),use_container_width=True,height=480)
        with tabs[3]:
            st.write("ساختار workbook اصلی، شیت‌ها، فرمول‌ها، استایل‌ها و سلول‌های دیگر با openpyxl نگه داشته می‌شوند؛ فقط قیمت‌های ویرایش‌شده نوشته می‌شوند.")
            source_name=st.session_state.raw_name or "products.xlsx"
            if source_name.lower().endswith(".xls"):
                st.warning("برای خروجی ویرایشی، فهرست کالا را به xlsx تبدیل کنید. گزارش خرید xls قابل قبول است.")
            else:
                try:
                    output=_to_xlsx(st.session_state.raw,source_name,st.session_state.sheet,st.session_state.edited,st.session_state.meta)
                    outname=Path(source_name).stem+"_قیمت_ویرایش_شده."+source_name.rsplit(".",1)[-1]
                    st.download_button("⬇️ دانلود اکسل با قیمت‌های ویرایش‌شده",data=output,file_name=outname,
                                       mime="application/vnd.ms-excel.sheet.macroEnabled.12" if outname.endswith("xlsm") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       type="primary",use_container_width=True)
                    st.caption(f"حجم خروجی: {len(output):,} بایت؛ تعداد ردیف‌های کالا: {len(st.session_state.edited):,}.")
                except Exception as exc:
                    st.error(f"ساخت فایل خروجی ناموفق بود: {exc}")

